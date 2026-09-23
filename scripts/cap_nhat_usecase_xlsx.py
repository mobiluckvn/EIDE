#!/usr/bin/env python3
"""Điền kết quả kiểm thử ngược vào `docs/Usecase_Test_Agent_Ky_Su_Nhung.xlsx`.

Sheet "Huong dan" nói rõ tester chỉ được điền năm cột nền vàng: *Trạng thái, Kết quả thực tế,
Người test, Ngày test, Ghi chú*. Script này chỉ chạm đúng năm cột ấy — đề bài (tên kịch bản,
các bước, kết quả mong đợi) là của chủ sản phẩm, không phải chỗ người chạy test sửa cho vừa
kết quả mình đo được.

Nguồn phán quyết là `docs/test/usecase/phan-quyet.json` — do NGƯỜI đọc nhật ký và ảnh chụp rồi
viết ra, không phải do bộ sàng dấu hiệu tự sinh. Bộ sàng chỉ biết "có nhắc tới X không".
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

GOC = Path(__file__).resolve().parent.parent
XLSX = GOC / "docs/Usecase_Test_Agent_Ky_Su_Nhung.xlsx"
PQ = GOC / "docs/test/usecase/phan-quyet.json"
RA = GOC / "docs/test/usecase"

HOP_LE = {"Chưa test", "Đạt", "Không đạt", "Bị chặn", "Bỏ qua"}


def main() -> int:
    import openpyxl

    pq = json.loads(PQ.read_text(encoding="utf-8")) if PQ.exists() else {}
    if not pq:
        print("chưa có phan-quyet.json — không điền gì", file=sys.stderr)
        return 1

    w = openpyxl.load_workbook(XLSX)
    ws = w["Test case"]
    hdr = [c.value for c in ws[1]]
    cot = {t: hdr.index(t) + 1 for t in
           ["Mã TC", "Trạng thái", "Kết quả thực tế", "Người test", "Ngày test", "Ghi chú"]}

    dem: dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        ma = ws.cell(r, cot["Mã TC"]).value
        if not ma or ma not in pq:
            continue
        p = pq[ma]
        tt = p["trang_thai"]
        if tt not in HOP_LE:
            print(f"  ⚠ {ma}: trạng thái '{tt}' không nằm trong danh sách cho phép",
                  file=sys.stderr)
        ws.cell(r, cot["Trạng thái"]).value = tt
        ws.cell(r, cot["Kết quả thực tế"]).value = p["ket_qua"]
        ws.cell(r, cot["Người test"]).value = "Claude Code (tự động qua EideApp --kich-ban)"
        ws.cell(r, cot["Ngày test"]).value = "23/09/2026"
        ws.cell(r, cot["Ghi chú"]).value = p.get("ghi_chu", "")
        dem[tt] = dem.get(tt, 0) + 1

    w.save(XLSX)
    tong = sum(dem.values())
    print(f"đã điền {tong}/76 test case vào {XLSX.name}")
    for k in sorted(dem, key=lambda x: -dem[x]):
        print(f"  {k:12} {dem[k]:3}")
    da_test = sum(v for k, v in dem.items() if k in {"Đạt", "Không đạt"})
    if da_test:
        print(f"  tỉ lệ đạt (trên {da_test} ca đã test thật): "
              f"{dem.get('Đạt', 0) / da_test:.0%}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
