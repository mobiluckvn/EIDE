#!/usr/bin/env python3
"""Sinh báo cáo tổng hợp `docs/test/usecase/BAO-CAO-TONG-HOP.docx` từ 76 phán quyết.

Mỗi TC đã có một `.docx` riêng để rà chi tiết. Tệp này trả lời câu hỏi khác: *nhìn tổng thể thì
sản phẩm đứng ở đâu* — đạt bao nhiêu, hỏng theo những NHÓM NGUYÊN NHÂN nào, và cái gì phải sửa
trước khi phát hành.

Xếp theo nguyên nhân chứ không theo mã TC: mười test case cùng trượt vì một gốc thì đó là MỘT
việc phải sửa, không phải mười. Đọc theo mã TC sẽ ra mười dòng rời rạc và một kế hoạch sửa sai
hình dạng.
"""

from __future__ import annotations

import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

GOC = Path(__file__).resolve().parent.parent
PQ = GOC / "docs/test/usecase/phan-quyet.json"
KE = GOC / "docs/test/usecase/ke-hoach.json"
DICH = GOC / "docs/test/usecase/BAO-CAO-TONG-HOP.docx"


def main() -> int:
    from docx import Document
    from docx.shared import Pt, RGBColor

    if not PQ.exists():
        print("chưa có phan-quyet.json", file=sys.stderr)
        return 1
    pq = json.loads(PQ.read_text(encoding="utf-8"))
    ke = {t["tc"]: t for t in json.loads(KE.read_text(encoding="utf-8"))}

    d = Document()
    d.add_heading("Kết quả kiểm thử bộ usecase Agent hỗ trợ kỹ sư nhúng", level=0)
    d.add_paragraph("Nguồn đề bài: docs/Usecase_Test_Agent_Ky_Su_Nhung.xlsx — 19 usecase, "
                    "76 test case. Ngày chạy: 23/09/2026. Sản phẩm: EIDE (apps/eide).")
    d.add_paragraph("Cách chạy: mỗi test case là một phiên thật, bộ lái gõ vào ô lệnh của vùng "
                    "trao đổi rồi bấm Gửi — không gọi tắt xuống năng lực. Mỗi bước chụp một ảnh "
                    "cửa sổ; cuối mỗi ca quét toàn bộ tab tác tử đã mở, mỗi tab một ảnh. Nhật ký "
                    "và ảnh của từng ca nằm trong tệp .docx cùng tên.")

    # ---- Bảng tổng
    dem = Counter(p["trang_thai"] for p in pq.values())
    theo_ut: dict[str, Counter] = defaultdict(Counter)
    for ma, p in pq.items():
        theo_ut[ke.get(ma, {}).get("uu_tien", "?")][p["trang_thai"]] += 1

    d.add_heading("1. Tổng quan", level=1)
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["Ưu tiên", "Đạt", "Không đạt", "Bị chặn", "Tổng"]):
        t.rows[0].cells[i].text = h
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for ut in ["P1", "P2", "P3"]:
        c = theo_ut.get(ut, Counter())
        r = t.add_row().cells
        r[0].text = ut
        r[1].text = str(c.get("Đạt", 0))
        r[2].text = str(c.get("Không đạt", 0))
        r[3].text = str(c.get("Bị chặn", 0))
        r[4].text = str(sum(c.values()))
    r = t.add_row().cells
    r[0].text = "TỔNG"
    r[1].text = str(dem.get("Đạt", 0))
    r[2].text = str(dem.get("Không đạt", 0))
    r[3].text = str(dem.get("Bị chặn", 0))
    r[4].text = str(sum(dem.values()))
    for c in r:
        c.paragraphs[0].runs[0].bold = True

    da = dem.get("Đạt", 0) + dem.get("Không đạt", 0)
    p = d.add_paragraph()
    p.add_run("Tỉ lệ đạt trên số ca test được thật: ").bold = True
    p.add_run(f"{dem.get('Đạt', 0)}/{da} = {dem.get('Đạt', 0) / da:.0%}" if da else "—")
    d.add_paragraph("Ca 'Bị chặn' không tính vào tỉ lệ: chúng chưa được đo, và gộp chúng vào "
                    "mẫu số là tự cho điểm một thứ chưa thi.")

    # ---- Nhóm nguyên nhân
    nhom: dict[str, list[str]] = defaultdict(list)
    for ma, p_ in pq.items():
        if p_["trang_thai"] == "Không đạt":
            nhom[p_.get("nhom", "khác")].append(ma)

    d.add_heading("2. Hỏng theo nhóm nguyên nhân", level=1)
    d.add_paragraph("Nhiều test case cùng trượt vì một gốc thì đó là MỘT việc phải sửa, không "
                    "phải nhiều. Xếp theo mã TC sẽ ra một kế hoạch sửa sai hình dạng.")
    for g, ds in sorted(nhom.items(), key=lambda x: -len(x[1])):
        d.add_heading(f"{g} — {len(ds)} test case", level=2)
        d.add_paragraph("Ca liên quan: " + ", ".join(sorted(ds)))

    # ---- Bảng chi tiết
    d.add_heading("3. Chi tiết từng test case", level=1)
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["Mã", "UC", "Ưu tiên", "Trạng thái", "Kết quả thực tế"]):
        t.rows[0].cells[i].text = h
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for ma in sorted(pq):
        p_ = pq[ma]
        k = ke.get(ma, {})
        r = t.add_row().cells
        r[0].text = ma
        r[1].text = k.get("uc", "")
        r[2].text = k.get("uu_tien", "")
        r[3].text = p_["trang_thai"]
        r[4].text = p_["ket_qua"]
        run = r[3].paragraphs[0].runs[0]
        run.font.color.rgb = {"Đạt": RGBColor(0x1B, 0x7F, 0x3B),
                              "Không đạt": RGBColor(0xC0, 0x28, 0x28),
                              "Bị chặn": RGBColor(0x94, 0x6C, 0x00)}.get(
            p_["trang_thai"], RGBColor(0x44, 0x44, 0x44))
        for c in r:
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8)

    d.save(str(DICH))
    print(f"đã sinh {DICH}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
