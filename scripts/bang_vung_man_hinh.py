"""Sinh `docs/EIDE-VUNG-MAN-HINH.xlsx` — bảng rà soát TỪNG VÙNG trên màn hình.

Chủ sản phẩm yêu cầu 22/09/2026: *"Vùng trao đổi dùng để trao đổi giữa người và Agent. Các
vùng còn lại theo menu thì bạn đề xuất cho tôi cụ thể danh sách — thiết kế thì show cái gì?
dữ liệu cần hiển thị là gì?"*

## Vì sao sinh bằng script chứ không gõ tay một lần

Ba cột trong bảng là **phép đo**, không phải ý kiến: `Hợp đồng nói gì` đọc từ
`docs/spec/ui/screens.json` (sinh từ `uxd.js`), `Hiện trạng gọi gì` dò từ chính mã Swift đang
chạy. Gõ tay thì hai cột ấy lệch khỏi kho ngay lần sửa mã kế tiếp, và một bảng rà soát nói sai
về hiện trạng thì tệ hơn không có — người đọc sẽ duyệt một thứ không tồn tại.

Cột `Đề xuất` là ý kiến, và nó được viết tay ở `DE_XUAT` bên dưới để chủ sản phẩm sửa thẳng.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

GOC = Path(__file__).resolve().parent.parent
SPEC = GOC / "docs" / "spec"
SWIFT = GOC / "apps" / "eide" / "Sources" / "EideGiaoDien"

# Nhóm điều hướng ở cột trái — lấy từ `EideManHinhDS.swift`, giữ đúng thứ tự người nhìn thấy.
NHOM = {
    "Main": "DỰ ÁN", "NhatKy": "DỰ ÁN", "FlowMap": "DỰ ÁN",
    "Ingest": "TRI THỨC", "Passport": "TRI THỨC", "Board": "TRI THỨC",
    "Graph": "TRI THỨC", "XungDot": "TRI THỨC",
    "LamRo": "THIẾT KẾ", "ReqArch": "THIẾT KẾ", "DiagramView": "THIẾT KẾ",
    "PlanDiff": "THIẾT KẾ", "Doc": "THIẾT KẾ",
    "Code": "MÃ NGUỒN", "DiffMerge": "MÃ NGUỒN",
    "Sim": "CHẠY THỬ", "Discovery": "CHẠY THỬ", "LogAssist": "CHẠY THỬ",
    "Debug": "CHẠY THỬ", "Bench": "CHẠY THỬ",
    "Env": "HỆ THỐNG", "Models": "HỆ THỐNG", "ToolForge": "HỆ THỐNG",
    "Registry": "HỆ THỐNG", "ChinhSach": "HỆ THỐNG",
}

# `so` trong screens.json → tiền tố màn trong mã. Hai bảng đánh số khác nhau: UXC-31 đếm cả
# Chat và khung vỏ, còn điều hướng chỉ đếm 25 màn mở được thành tab.
THEO_SO = {
    "2": "Main", "4": "Ingest", "5": "Passport", "6": "Board", "7": "Graph",
    "8": "NhatKy", "9": "XungDot", "10": "LamRo", "11": "ReqArch", "12": "DiagramView",
    "13": "Doc", "14": "PlanDiff", "15": "Code", "16": "Sim", "17": "Discovery",
    "18": "LogAssist", "19": "Debug", "20": "ToolForge", "21": "Bench", "22": "Registry",
    "23": "Models", "24": "Env", "25": "FlowMap", "26": "DiffMerge", "27": "ChinhSach",
}

# ---- Ý KIẾN, sửa thẳng ở đây. (mục đích một câu, dữ liệu PHẢI hiện, hành động trên màn)
DE_XUAT: dict[str, tuple[str, str, str]] = {
 "Main": (
  "Trả lời trong 5 giây: dự án đang ở đâu, việc gì đang chờ TÔI.",
  "• Tính năng: n passing / n tổng, cái đang làm dở\n"
  "• Đang chờ tôi: số mục + 3 cái gần nhất, bấm được sang đúng màn\n"
  "• Hoàn tác được: số mục + hạn sớm nhất\n"
  "• Chip/board đã ghim + ISA + trạng thái hộ chiếu\n"
  "• Chi phí mô hình hôm nay / hạn ngày\n"
  "• 3 việc tác tử vừa làm xong (có giờ)",
  "Bấm một dòng → mở đúng màn của việc ấy"),
 "NhatKy": (
  "Mọi việc đã xảy ra, truy được — đây là sổ cái, không phải log kỹ thuật.",
  "• Thời điểm · loại · năng lực · kết quả · chi phí\n"
  "• Lọc: chỉ việc của tôi / chỉ việc tác tử tự làm / chỉ lỗi / chỉ cổng\n"
  "• Mỗi dòng nói bằng tiếng Việt việc gì đã xảy ra, không dán tên sự kiện",
  "Bấm một dòng → mở màn tương ứng; Xuất nhật ký"),
 "FlowMap": (
  "Tôi đang ở bước nào của quy trình, bước kế tiếp là gì.",
  "• Bản đồ P0–P7, tô màu bước đã qua / đang làm / chưa tới\n"
  "• Mỗi bước: cổng nào canh, quyết định gần nhất của cổng ấy + lý do\n"
  "• Bước đang bị chặn hiện nổi bật kèm CÁCH GỠ",
  "Bấm một bước → mở màn làm bước ấy"),
 "Ingest": (
  "Tri thức vào dự án từ đâu, và đã rút ra được gì.",
  "• Danh mục NGUỒN: tên · loại · tầng tin cậy · số fact rút được · ngày\n"
  "• Lượt nhập gần nhất: tiến độ, cái nào hỏng và vì sao\n"
  "• Tài liệu tác tử TỰ tải về (nói rõ từ đâu, giấy phép gì)",
  "Kéo thả tệp/zip; Dán URL; Xoá nguồn (kèm cảnh báo fact nào mất theo)"),
 "Passport": (
  "Con chip này, máy biết CHẮC những gì — và tin được tới đâu.",
  "• Bộ nhớ (FLASH/RAM/EEPROM), ngoại vi, chân, thanh ghi\n"
  "• Mỗi con số kèm: tầng (vàng/bạc/đồng), trạng thái duyệt, NGUỒN + số trang\n"
  "• Số fact chưa duyệt — bấm được sang duyệt\n"
  "• Phiên bản hộ chiếu đã ghim cho dự án",
  "Xem nguồn (mở PDF đúng trang); Duyệt fact; Ghim phiên bản khác"),
 "Board": (
  "Bo mạch thật: chân nào nối đi đâu, có đụng nhau không.",
  "• Bảng net/pin, nguồn (netlist/BOM/ảnh sơ đồ)\n"
  "• XUNG ĐỘT chân hiện trên cùng, kèm đề xuất sửa\n"
  "• Board đã đánh dấu 'lab' hay chưa (quyết định R3)",
  "Nhập netlist/BOM; Đánh dấu board lab; Nhận đề xuất sửa chân"),
 "Graph": (
  "Hỏi bằng tiếng Việt về tri thức dự án, nhận câu trả lời CÓ TRÍCH DẪN.",
  "• Ô hỏi + câu trả lời kèm fact id và nguồn\n"
  "• Bản đồ lân cận của thứ đang hỏi, tô màu theo tầng/trạng thái\n"
  "• Đường truy nguồn: fact → tài liệu → trang",
  "Hỏi; Bấm một nút → xem fact; Mở nguồn"),
 "XungDot": (
  "Hai nguồn nói khác nhau về cùng một thứ — tôi chọn bên nào.",
  "• Hai bên ĐẶT CÙNG HÀNG: giá trị, nguồn, tầng, độ tin, ngày\n"
  "• Vì sao máy cho là mâu thuẫn (suy ra hay được khai)\n"
  "• Hệ quả: mã nào đang dùng con số này",
  "Chọn A / Chọn B / Cả hai sai — ghi lý do"),
 "LamRo": (
  "Tác tử đang chờ tôi cho biết gì — VÀ tôi trả lời được ngay ở đây.",
  "• Từng câu hỏi: hỏi gì, vì bước nào, lựa chọn có sẵn (nếu là tập đóng)\n"
  "• Ô trả lời ngay trên dòng; lịch sử trả lời + nút hoàn tác từng lần\n"
  "• Câu đã trả lời xếp xuống dưới, không mất",
  "Trả lời; Hoàn tác câu trả lời; Bỏ qua (ghi lý do)"),
 "ReqArch": (
  "Tác tử hiểu việc của tôi thành những yêu cầu nào.",
  "• ReqSet: mã · loại · ưu tiên · nội dung · tiêu chí nghiệm thu\n"
  "• Yêu cầu KHÔNG đo được đánh dấu riêng\n"
  "• Ma trận truy vết yêu cầu ↔ module ↔ mã ↔ test\n"
  "• ADR: quyết định kiến trúc + lý do + phương án đã loại",
  "Sửa một yêu cầu; Thêm tiêu chí; Duyệt tập yêu cầu"),
 "DiagramView": (
  "Nhìn thấy hệ thống, và hình luôn khớp mã.",
  "• Hình + mã nguồn của hình cạnh nhau\n"
  "• Cờ LỆCH khi mã đổi mà hình chưa sinh lại\n"
  "• Lược đồ nào đang dùng trong tài liệu nào",
  "Sinh lại; Đồng bộ; Chèn vào tài liệu; Xuất PNG/SVG"),
 "PlanDiff": (
  "Tác tử định làm gì, theo thứ tự nào, dựa trên căn cứ gì.",
  "• Khối CÒN THIẾU đầu bảng, nói rõ thiếu gì và ai trả lời được\n"
  "• Từng bước: mục tiêu · năng lực · fact trích dẫn · điều kiện xong\n"
  "• Quyết định cổng G1 + nút DUYỆT ngay trên màn\n"
  "• Ước lượng chi phí so với hạn ngày",
  "Duyệt kế hoạch; Từ chối kèm lý do; Lập lại"),
 "Doc": (
  "Bộ tài liệu gửi được cho người khác.",
  "• Danh sách tài liệu: loại · phiên bản · ngày sinh · mục nào STALE\n"
  "• Xem trước nội dung; lỗi style check\n"
  "• Lược đồ đã nhúng",
  "Sinh; Xuất PDF/DOCX; Sinh lại mục stale"),
 "Code": (
  "Mã tác tử viết, và mọi hằng số truy được về datasheet.",
  "• Cây tệp + nội dung; hằng số tô màu, hover ra fact + nguồn\n"
  "• Vi phạm constant-guard hiện đỏ tại dòng\n"
  "• Tệp nào do tác tử viết, tệp nào tôi sửa tay",
  "Sửa và Lưu; Gửi cổng merge; Hoàn tác lần lưu"),
 "DiffMerge": (
  "Trước khi mã vào dự án, bốn cổng nói gì.",
  "• Diff hai cột, theo tệp\n"
  "• G1/G3/G4/G5: đạt hay chặn, LÝ DO chặn, cần gì để qua\n"
  "• Kết quả 4 công cụ kiểm + reviewer khác hãng",
  "Duyệt merge; Từ chối; Yêu cầu sửa"),
 "Sim": (
  "Chạy thử firmware TRƯỚC khi đụng phần cứng.",
  "• Nền tảng mô phỏng: engine, chip, bản đồ bộ nhớ (kèm fact trích dẫn)\n"
  "• Kịch bản + kết quả từng lần chạy: đạt/hỏng, số đo, log\n"
  "• Ngoại vi nào mô phỏng được, cái nào KHÔNG (nói thẳng)",
  "Chạy; Dựng lại nền tảng; Mở log; So hai lần chạy"),
 "Discovery": (
  "Cắm board vào thì máy thấy gì.",
  "• Cổng/probe tìm thấy, ID chip đọc được, tốc độ\n"
  "• Đối chiếu với chip đã ghim: KHỚP hay LỆCH\n"
  "• Bus quét được, nguồn điện",
  "Dò lại; Ghi vào target.yaml; Đánh dấu board lab"),
 "LogAssist": (
  "Log dài, tôi hỏi tại một dòng và được giải thích.",
  "• Log có lọc theo mẫu lỗi; serial trực tiếp\n"
  "• Hỏi tại dòng → trả lời kèm fact/mã liên quan",
  "Lọc; Hỏi tại dòng; Lưu đoạn log làm bằng chứng"),
 "Debug": (
  "Lỗi này do đâu — có bằng chứng, không đoán.",
  "• Giả thuyết đang xét + bằng chứng ủng hộ/phản bác\n"
  "• EvidencePack: log, số đo, thanh ghi đọc được\n"
  "• Thí nghiệm tiếp theo tác tử đề nghị",
  "Chạy thí nghiệm; Loại giả thuyết; Kết luận"),
 "Bench": (
  "Đo sản phẩm bằng bài chuẩn, không bằng cảm giác.",
  "• Bài CF/BF/BC: điểm, thời gian, chi phí\n"
  "• Huy hiệu đạt được; so với lần chạy trước",
  "Chạy bài; So hai lần"),
 "Env": (
  "Máy tôi thiếu gì để làm việc này.",
  "• Từng công cụ: có/không, phiên bản, đường dẫn\n"
  "• ISA của dự án cần gì, thiếu cái nào hiện ĐỎ\n"
  "• Lệnh cài CHẠY ĐƯỢC cho đúng máy này",
  "Cài tự động; Chép lệnh cài; Khoá phiên bản"),
 "Models": (
  "Tiền đang tiêu vào đâu.",
  "• Vai trò → mô hình đang dùng\n"
  "• Chi phí hôm nay / hạn ngày, theo vai trò\n"
  "• Lượt gọi gần nhất: vai trò, token, giá",
  "Đổi mô hình cho một vai trò; Đặt hạn ngày"),
 "ToolForge": (
  "Tác tử tự viết công cụ gì cho dự án này.",
  "• Danh sách công cụ tự tạo: làm gì, đã test chưa, dùng bao nhiêu lần\n"
  "• Mã công cụ + kết quả test",
  "Xem mã; Chạy test; Thăng cấp thành năng lực"),
 "Registry": (
  "Gói tri thức dùng lại được giữa các dự án.",
  "• Gói có sẵn: hộ chiếu chip, mẫu dự án, skill\n"
  "• Chữ ký + giấy phép từng gói",
  "Tải về; Đóng gói dự án này; Xuất bản"),
 "ChinhSach": (
  "Tác tử được phép tự làm tới đâu.",
  "• Mức tự chủ hiện tại + nghĩa của nó bằng tiếng Việt\n"
  "• Bảng quy tắc ĐANG CÓ HIỆU LỰC, quy tắc nào vừa chặn việc gì\n"
  "• Trạng thái niêm phong chính sách",
  "Đổi mức tự chủ; Xem lịch sử quyết định của một quy tắc"),
}


# ---- Ý KIẾN: phải sửa gì để đạt cột "DỮ LIỆU PHẢI HIỆN". Viết tay, sửa thẳng ở đây.
#
# Ba loại việc, cố ý tách vì chúng đi ba đường khác nhau:
#   [GD]  chỉ giao diện — dữ liệu đã có trong store/năng lực, màn chưa hiện
#   [NL]  cần năng lực mới hoặc sửa năng lực đang có
#   [HC]  chờ phần cứng thật hoặc năng lực chạm phần cứng (22 năng lực còn thiếu)
CAN_SUA: dict[str, str] = {
 "Main": "[GD] Màn đang đọc `budget.state` + `session.state`; thiếu hẳn khối 'đang chờ tôi' "
         "và 'hoàn tác được' — hai thứ ấy đã có sẵn ở `queue.list` và `undo.list`, chỉ chưa "
         "gọi.\n[GD] Chip/ISA/hộ chiếu: `project.status` trả rồi, chưa hiện.\n"
         "[GD] Bấm một dòng → mở đúng màn: chưa nối.",
 "NhatKy": "[GD] Đang hiện `view.timeline` thô, mỗi dòng là tên sự kiện sổ cái. Cần dịch sang "
           "tiếng Việt theo `kind` và thêm bộ lọc (của tôi / tác tử tự làm / lỗi / cổng).\n"
           "[GD] Cột chi phí: `model.call` có `cost_usd`, chưa gộp vào dòng.",
 "FlowMap": "[GD] Đang đọc `view.timeline`; chưa vẽ bản đồ P0–P7 và chưa gắn quyết định cổng "
            "vào từng bước — `decision_log` đã có đủ (gate, rule, reason, action_cap).\n"
            "[GD] Bước đang bị chặn + CÁCH GỠ: `clarification` đã có sau [DEV-178], chưa nối.",
 "Ingest": "[GD] `archive.sources` đã có (DEV-134) và màn đã gọi — kiểm lại xem có hiện đủ "
           "tier/số fact/ngày không.\n[GD] Lượt nhập hỏng vì sao: `cap.run.finish` có error, "
           "chưa hiện.",
 "Passport": "[GD] `passport.query` trả fact + citations; màn mới hiện một phần. Thiếu: số "
             "fact CHƯA duyệt (đếm `status='normalized'`), phiên bản đã ghim "
             "(`constraints.target.pins.chip`).\n[GD] Nút 'Duyệt fact' → `kg.review_facts` "
             "chưa nối.",
 "Board": "[GD] `board.check_pins`/`board.constraints`/`board.propose_fix` đã gọi. Thiếu cờ "
          "'board lab' — nằm ở `autonomy.yaml/boards`, chưa hiện.\n"
          "[NL] Nhập netlist/BOM từ màn: đi qua `extract.kicad_netlist`/`extract.bom`, chưa "
          "có nút.",
 "Graph": "[GD] `view.rag_ask`/`view.rag_trace` đã gọi. Bản đồ lân cận (`kg.neighborhood`) và "
          "đường truy nguồn (`view.provenance`) đã có năng lực, màn chưa vẽ.",
 "XungDot": "[GD] `kg.conflicts` đã gọi. Thiếu cột 'hệ quả: mã nào đang dùng con số này' — "
            "`kg.impact` đã có, chưa nối.\n[GD] Nút chọn A/B → `kg.resolve_conflict`: kiểm "
            "lại đã nối chưa.",
 "LamRo": "[GD] ƯU TIÊN CAO — ô trả lời phải có CẢ ở vùng trao đổi (thẻ hỏi–đáp, khuôn "
          "`EideTheYHieu`), không chỉ ở tab này.\n"
          "[NL] Câu hỏi đang dán TÊN TRƯỜNG hợp đồng (`isa`). Cần: (a) đọc `description` của "
          "`input_schema`; (b) tham số có tập giá trị đóng thì liệt kê lựa chọn — `isa` lấy "
          "từ `docs/spec/isa/*.yaml`, `chip` từ bảng `passport`.\n"
          "[GD] Lịch sử trả lời + hoàn tác từng lần: `clarification_answer` đã có "
          "([DEV-151]), màn chưa hiện.",
 "ReqArch": "[GD] Đang đọc `requirement` + `adr`. Thiếu: đánh dấu yêu cầu KHÔNG đo được "
            "(`req.detect_conflict.issues` đã tính, đang chỉ vào tab Làm rõ), ma trận truy "
            "vết (`req.trace_matrix` đã có năng lực), HwMap + ngân sách RAM/Flash "
            "(`arch.map_hw`, `arch.budget`).",
 "DiagramView": "[GD] `diagram.sync` đã gọi, cờ LỆCH đã có. Thiếu: lược đồ nào đang dùng "
                "trong tài liệu nào (`doc_artifact` có liên kết, chưa hiện).",
 "PlanDiff": "[GD] Khối 'còn thiếu' đã đúng sau [DEV-171]/[DEV-174a].\n"
             "[GD] ƯU TIÊN CAO — nút DUYỆT kế hoạch: lõi đã xong ở [DEV-176] "
             "(`gate.decide` nhận khoá `<run_id>:<gate>`), GIAO DIỆN CHƯA CÓ NÚT.\n"
             "[GD] Fact trích dẫn của từng bước: `plan.steps[].cites` đã có sau [DEV-175], "
             "chưa hiện thành liên kết bấm được.",
 "Doc": "[GD] Đang đọc hiện vật `doc`. Thiếu: mục STALE (`doc_artifact.stale_sections` đã "
        "có), lỗi style (`doc.style_check`), nút Xuất (`report.export`).",
 "Code": "[GD] `code.constant_guard` đã gọi. Thiếu: hover hằng số ra fact + nguồn — "
         "`view.provenance` đã có; phân biệt tệp tác tử viết / tôi sửa tay (`code_unit` có "
         "`run_id` và `human.file_save` trong sổ cái).",
 "DiffMerge": "[GD] Thiếu bảng bốn cổng G1/G3/G4/G5 kèm LÝ DO chặn — `decision_log` đã có đủ "
              "sau [DEV-171]; đây là màn hưởng lợi trực tiếp.\n"
              "[GD] Kết quả 4 công cụ kiểm: `tool_report` đã có, chưa hiện.",
 "Sim": "[GD] Nền tảng mô phỏng + bản đồ bộ nhớ + trích dẫn fact: `sim.build_platform` trả "
        "đủ sau [DEV-173], màn chưa hiện.\n[GD] Ngoại vi KHÔNG mô phỏng được: `coverage."
        "unsupported` đã có, phải nói thẳng chứ không im.\n[HC] `sim.compare_hil` chưa hiện "
        "thực — phần so sánh với board thật chờ.",
 "Discovery": "[GD] MÀN CHƯA DỰNG — cần lớp `EideManDoBoard` và đăng ký vào `EidePhien.MAN`.\n"
              "[HC] 8/12 năng lực `discover.*` chưa hiện thực (probe, chip_id, bus_scan, "
              "link_speed, power, clock_measure, firmware_probe, board_match). Không có "
              "chúng thì màn chỉ là cái vỏ.",
 "LogAssist": "[GD] MÀN CHƯA DỰNG.\n[NL] `debug.log_stats`/`debug.ask_at` ĐÃ hiện thực — "
              "phần log lớn + hỏi tại dòng làm được NGAY, không chờ phần cứng.\n"
              "[HC] Riêng serial trực tiếp chờ `target.serial`.",
 "Debug": "[GD] MÀN CHƯA DỰNG.\n[NL] `debug.*` đã hiện thực đủ (6/6) — EvidencePack, giả "
          "thuyết, thí nghiệm làm được ngay.\n[HC] Đọc/ghi thanh ghi qua probe chờ "
          "`target.probe_read`/`probe_write`.",
 "Bench": "[GD] MÀN CHƯA DỰNG.\n[HC] `bench.run` chưa hiện thực; `measure.*` 0/3. Đây là màn "
          "phụ thuộc phần cứng nhiều nhất — đề nghị hoãn sau cùng.",
 "Env": "[GD] `env.check` đã gọi. Thiếu: đánh ĐỎ công cụ mà ISA của dự án cần — bảng "
        "`docs/spec/isa/*.yaml` đã có `toolchain.tools`.\n[GD] Lệnh cài chạy được: đã đúng "
        "sau [DEV-172], hiện ra màn là xong.",
 "Models": "[GD] `budget.state` đã gọi. Thiếu: vai trò → mô hình (`models.yaml`), chi phí "
           "theo vai trò (`model.call` trong sổ cái có `role` + `cost_usd`).",
 "ToolForge": "[GD] Đang đọc `view.artifacts(tool)`. Thiếu: mã công cụ + kết quả test "
              "(`tool_report` đã có), nút thăng cấp (`tool.promote`).",
 "Registry": "[GD] `registry.search` đã gọi. Thiếu: chữ ký + giấy phép từng gói "
             "(`registry.pack` có, `manifest` đã mang đủ).",
 "ChinhSach": "[GD] `policy.rules` đã gọi. Thiếu: quy tắc nào VỪA chặn việc gì — "
              "`decision_log` nay đủ dữ liệu sau [DEV-171].\n[GD] Nghĩa của mức tự chủ bằng "
              "tiếng Việt: đang hiện mã A0–A4 trần.",
}

def nang_luc_thieu() -> dict[str, bool]:
    """`id năng lực → đã hiện thực chưa`, hỏi CHÍNH registry mà sản phẩm nạp lúc chạy.

    Không đọc cột `m0`/`milestone` của cds.json: hai cột ấy nói KẾ HOẠCH, còn câu hỏi ở đây là
    "hôm nay gọi được chưa". Đo 22/09/2026: 223/245 hiện thực, 22 còn thiếu và gần như tất cả
    đều chạm phần cứng thật.
    """
    import sys
    sys.path.insert(0, str(GOC / "src"))
    from eide_core.registry import get_registry
    r = get_registry()
    ds = json.loads((SPEC / "cds.json").read_text(encoding="utf-8"))
    caps = ds["capabilities"] if isinstance(ds, dict) else ds
    return {c["id"]: (c["id"] in r and r.get(c["id"]).implemented) for c in caps}


def thieu_cua_man(mau: list[str], ht: dict[str, bool]) -> list[str]:
    """Năng lực hợp đồng giao cho màn này mà HÔM NAY chưa gọi được."""
    import fnmatch
    ra = set()
    for m in mau:
        m = m.strip()
        for cid, ok in ht.items():
            if not ok and (fnmatch.fnmatch(cid, m) if "*" in m else cid == m):
                ra.add(cid)
    return sorted(ra)


def hien_trang() -> dict[str, str]:
    """Dò từ CHÍNH mã Swift đang chạy — cột này là phép đo, không phải ý kiến."""
    ra: dict[str, str] = {}
    for f in sorted(SWIFT.glob("Eide*.swift")):
        s = f.read_text(encoding="utf-8")
        for m in re.finditer(
                r"class\s+(\w+)\s*:\s*EideMan\w*\s*\{(.*?)"
                r"(?=\n(?:public (?:final )?|open |final )?class |\Z)", s, re.S):
            than = m.group(2)
            t = re.search(r'var tien: String \{ "(\w+)" \}', than)
            if not t:
                continue
            goi = sorted(set(re.findall(r'goi\(\s*"([\w.]+)"', than))
                         | set(re.findall(r'nangLuc\(goi,\s*"([\w.]+)"', than)))
            hv = sorted(set(re.findall(r'hienVat\(goi,\s*"([\w.]+)"', than)))
            phan = []
            if hv:
                phan.append("hiện vật: " + ", ".join(hv))
            goi = [x for x in goi if x != "caps.invoke"]
            if goi:
                phan.append("gọi: " + ", ".join(goi))
            if not phan:
                # Vài màn không đi qua hai hàm tiện ích trên mà dùng thẳng `goi("…")` gói
                # trong một hàm riêng, hoặc đọc từ sự kiện đẩy. Quét MỌI chuỗi có dấu chấm
                # trông như tên năng lực / sự kiện — rộng hơn nên ghi rõ là suy đoán, để
                # người đọc biết cột này chắc tới đâu.
                rong_hon = sorted({x for x in re.findall(r'"([a-z][\w]*\.[\w.]+)"', than)
                                   if not x.endswith((".swift", ".json", ".md", ".png"))})
                if rong_hon:
                    phan.append("(dò rộng) " + ", ".join(rong_hon[:8]))
            ra[t.group(1)] = "; ".join(phan) or "KHÔNG dò được nguồn dữ liệu — cần xem tay"
    return ra


def main() -> int:
    ds = json.loads((SPEC / "ui" / "screens.json").read_text(encoding="utf-8"))
    ht = hien_trang()
    da_lam = nang_luc_thieu()

    wb = Workbook()
    ws = wb.active
    ws.title = "Vùng màn hình"
    cot = ["Nhóm", "Mã", "Màn", "Mục đích (một câu)", "DỮ LIỆU PHẢI HIỆN",
           "Hành động trên màn", "CẦN SỬA GÌ ĐỂ ĐẠT", "Năng lực CHƯA hiện thực",
           "HIỆN TRẠNG trong mã", "Hợp đồng UXC-31 nói gì", "Năng lực theo hợp đồng",
           "Ý kiến của anh"]
    ws.append(cot)
    dam = Font(bold=True, color="FFFFFF")
    nen = PatternFill("solid", fgColor="2F5496")
    for i, _ in enumerate(cot, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = dam, nen
        c.alignment = Alignment(vertical="center", wrap_text=True)

    thu_tu = ["DỰ ÁN", "TRI THỨC", "THIẾT KẾ", "MÃ NGUỒN", "CHẠY THỬ", "HỆ THỐNG"]
    hang = []
    for s in ds:
        tien = THEO_SO.get(str(s.get("so")))
        if not tien:
            continue                      # Chat và khung vỏ: đã chốt, không đưa vào bảng rà
        d = DE_XUAT.get(tien, ("", "", ""))
        thieu = thieu_cua_man(s["nang_luc"], da_lam)
        hang.append([NHOM.get(tien, "?"), f"S{s['so']}", s["man_hinh"], d[0], d[1], d[2],
                     CAN_SUA.get(tien, ""),
                     (", ".join(thieu) if thieu else "— đủ cả"),
                     ht.get(tien, "CHƯA DỰNG — có trong menu cột trái, KHÔNG có lớp màn trong EidePhien.MAN; bấm vào không mở gì"),
                     s["noi_dung"], ", ".join(s["nang_luc"]), ""])
    hang.sort(key=lambda r: (thu_tu.index(r[0]) if r[0] in thu_tu else 9, r[2]))
    for r in hang:
        ws.append(r)

    rong = [12, 6, 26, 40, 52, 34, 62, 30, 40, 44, 28, 26]
    for i, w in enumerate(rong, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, ws.max_row + 1):
        for i in range(1, len(cot) + 1):
            ws.cell(row=r, column=i).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "D2"

    # Trang 2: chốt về vùng trao đổi — chủ sản phẩm đã quyết, ghi lại để không bàn lại.
    w2 = wb.create_sheet("Vùng trao đổi")
    w2.append(["Nguyên tắc", "Chi tiết"])
    for i in (1, 2):
        c = w2.cell(row=1, column=i)
        c.font, c.fill = dam, nen
    for a, b in [
        ("Vai trò", "Nơi DUY NHẤT người và tác tử trao đổi. Mọi câu tác tử hỏi phải hiện ở "
                    "đây, và phải trả lời được NGAY tại đây."),
        ("Quan hệ với tab Làm rõ yêu cầu",
         "Vùng trao đổi là CUỘC TRÒ CHUYỆN (hỏi và trả lời). Tab là SỔ GHI (còn treo những "
         "gì). Cùng một câu hỏi xuất hiện ở cả hai; trả lời ở đâu cũng được."),
        ("Câu hỏi phải viết thế nào",
         "Tiếng Việt, không dán tên trường hợp đồng. Tham số có tập giá trị đóng thì LIỆT KÊ "
         "lựa chọn kèm ví dụ (ví dụ `isa`: armv7e-m / avr8 / rv32imac)."),
        ("Thiếu ở bản 22/09/2026",
         "Vùng trao đổi hiện đúng một dòng 'Dừng ở `env.check` — cần anh cho biết: isa': dán "
         "tên trường, không nói lựa chọn, KHÔNG có ô trả lời. Người dùng phải đi tìm tab."),
    ]:
        w2.append([a, b])
    w2.column_dimensions["A"].width = 30
    w2.column_dimensions["B"].width = 100
    for r in range(2, w2.max_row + 1):
        for i in (1, 2):
            w2.cell(row=r, column=i).alignment = Alignment(vertical="top", wrap_text=True)

    ra = GOC / "docs" / "EIDE-VUNG-MAN-HINH.xlsx"
    wb.save(ra)
    print(f"đã sinh {ra} — {len(hang)} màn")
    return 0


if __name__ == "__main__":
    sys.exit(main())
