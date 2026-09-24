#!/usr/bin/env python3
"""Chạy bộ kiểm thử `docs/Usecase_Test_Agent_Ky_Su_Nhung.xlsx` qua ĐÚNG đường giao diện.

Mỗi test case là một phiên thật: bộ lái `EideApp --kich-ban` gõ vào ô lệnh, bấm Gửi, đợi tác tử
trả lời, chụp màn hình từng bước và ghi nhật ký. Không gọi tắt xuống `chat.send` — cái đó đo
được năng lực nhưng KHÔNG đo được thứ bộ kiểm thử này hỏi: *người ngồi trước máy có nhận được
câu trả lời đúng không*.

## Vì sao mỗi TC một phiên riêng

Chạy chung một phiên thì lượt sau thừa hưởng ngữ cảnh lượt trước, và một TC "đạt" nhờ câu trả
lời của TC khác là một phép đo nói dối. Vài TC CỐ Ý nối tiếp nhau (TC002 nối TC001) — những TC
ấy khai `noi_tiep` và dùng lại thư mục dự án của TC trước, mọi TC khác bắt đầu sạch.

## Dấu hiệu máy đọc được, phán quyết do người

`dau_hieu.phai_co` / `khong_duoc_co` chỉ là phép sàng: chúng bắt được "có nhắc tới X không",
không bắt được "câu trả lời có ĐÚNG không". Máy chấm xong ra `Cần người đọc`; người đọc nhật ký
và ảnh rồi mới chốt Đạt/Không đạt. Một bộ kiểm thử tự chấm mình toàn màu xanh là bộ kiểm thử
không ai tin được.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

import thu_toan_canh

GOC = Path(__file__).resolve().parent.parent
APP = GOC / "apps/eide/.build/debug/EideApp"
KE_HOACH = GOC / "docs/test/usecase/ke-hoach.json"
RA = GOC / "docs/test/usecase"

HAN_GIAY = 900


def _chay_kich_ban(tc: dict, thu_muc: Path, du_an_truoc: Path | None) -> dict:
    """Chạy một TC qua bộ lái kịch bản. Trả về nhật ký, ảnh, thời gian, mã thoát."""
    thu_muc.mkdir(parents=True, exist_ok=True)
    buoc = list(tc["buoc"])
    if tc.get("noi_tiep") and du_an_truoc is not None:
        buoc.insert(0, f"@mo {du_an_truoc}")

    kb = thu_muc / "kich-ban.kb"
    kb.write_text(
        f"# {tc['tc']} — {tc['ten']}\n" + "\n".join(buoc) + "\n", encoding="utf-8"
    )

    t0 = time.time()
    p = subprocess.run(
        [str(APP), "--kich-ban", str(kb), "--ra", str(thu_muc)],
        capture_output=True,
        text=True,
        timeout=HAN_GIAY * len(buoc) + 120,
        cwd=GOC,
    )
    giay = time.time() - t0

    (thu_muc / "stdout.log").write_text(p.stdout + "\n--- stderr ---\n" + p.stderr,
                                        encoding="utf-8")
    nhat = (thu_muc / "nhat-ky.md")
    return {
        "giay": round(giay, 1),
        "rc": p.returncode,
        "nhat_ky": nhat.read_text(encoding="utf-8") if nhat.exists() else p.stdout,
        # CẢ ảnh từng bước LẪN ảnh từng tab của lượt quét — docx phải mang đủ bằng chứng,
        # và phần lớn thứ tác tử làm ra nằm ở các tab, không ở khung chat.
        "anh": (sorted(x.name for x in thu_muc.glob("buoc-*.png"))
                + sorted(x.name for x in thu_muc.glob("man-*.png"))),
        "du_an": next(iter((thu_muc / "du-an").glob("*")), None) if (thu_muc / "du-an").is_dir()
        else None,
    }


def _cham(tc: dict, nhat_ky: str) -> dict:
    """Sàng bằng dấu hiệu. KHÔNG kết luận Đạt — chỉ nói dấu hiệu nào thiếu."""
    d = tc.get("dau_hieu", {})
    van = nhat_ky.lower()

    def co(m: str) -> bool:
        # Mỗi dấu hiệu có thể là nhiều cách nói cùng một ý, phân cách bằng '|'.
        return any(re.search(x.strip().lower(), van) for x in m.split("|"))

    thieu = [m for m in d.get("phai_co", []) if not co(m)]
    thua = [m for m in d.get("khong_duoc_co", []) if co(m)]
    return {"thieu": thieu, "thua": thua,
            "so_bo": "Cần người đọc" if not thieu and not thua else "Nghi Không đạt"}


def _docx(tc: dict, kq: dict, cham: dict, thu_muc: Path, dich: Path) -> None:
    """Một TC = một tệp .docx để chủ sản phẩm rà lại: đề bài, việc đã làm, bằng chứng."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Pt(50)

    h = d.add_heading(f"{tc['tc']} — {tc['ten']}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in [
        ("Mã usecase", tc["uc"]),
        ("Loại", tc["loai"]),
        ("Ưu tiên", tc["uu_tien"]),
        ("Chế độ chạy", tc.get("che_do", "ui")),
        ("Thời gian chạy", f"{kq['giay']} s"),
        ("Người test", "Claude Code (tự động, qua EideApp --kich-ban)"),
        ("Ngày test", tc.get("ngay", "")),
        ("Trạng thái", tc.get("trang_thai", cham["so_bo"])),
        ("Số lần chạy", tc.get("ti_le_qua_sang", "1 lần")),
    ]:
        r = t.add_row().cells
        r[0].text = str(k)
        r[1].text = str(v)
        r[0].paragraphs[0].runs[0].bold = True

    d.add_heading("1. Đề bài (nguyên văn từ bảng test case)", level=2)
    for nhan, khoa in [("Tiền điều kiện", "tien_dk"), ("Các bước / Đầu vào", "vao"),
                       ("Kết quả mong đợi", "cho")]:
        p = d.add_paragraph()
        p.add_run(f"{nhan}: ").bold = True
        p.add_run(str(tc.get(khoa, "") or "—"))

    d.add_heading("2. Đã làm gì", level=2)
    d.add_paragraph(tc.get("cach_lam", "Chạy qua bộ lái kịch bản của EideApp: gõ vào ô lệnh "
                                       "của vùng trao đổi, bấm Gửi, đợi tác tử trả lời xong."))
    d.add_paragraph("Các dòng kịch bản đã chạy:")
    for b in tc["buoc"]:
        d.add_paragraph(b, style="List Bullet")

    d.add_heading("3. Kết quả thực tế", level=2)
    p = d.add_paragraph()
    p.add_run("Phán quyết: ").bold = True
    r = p.add_run(tc.get("trang_thai", cham["so_bo"]))
    r.bold = True
    r.font.color.rgb = {
        "Đạt": RGBColor(0x1B, 0x7F, 0x3B), "Không đạt": RGBColor(0xC0, 0x28, 0x28),
        "Bị chặn": RGBColor(0x94, 0x6C, 0x00),
    }.get(tc.get("trang_thai", ""), RGBColor(0x44, 0x44, 0x44))
    d.add_paragraph(tc.get("nhan_xet", "(chưa có nhận xét của người đọc)"))

    if kq.get("lap") and len(kq["lap"]) > 1:
        d.add_paragraph(
            f"Chạy lặp {len(kq['lap'])} lần (sheet \"Huong dan\" đòi 3–5 lần vì kết quả của "
            f"mô hình ngôn ngữ không tất định). Ảnh và nhật ký dưới đây là của LẦN ĐẦU — lấy "
            f"\"lần đẹp nhất\" thì bộ đo thành bộ chọn kết quả.")
        tl = d.add_table(rows=1, cols=4)
        tl.style = "Table Grid"
        for i, h in enumerate(["Lần", "Thời gian", "Qua sàng", "Dấu hiệu thiếu"]):
            tl.rows[0].cells[i].text = h
            tl.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        for i, x in enumerate(kq["lap"], 1):
            r2 = tl.add_row().cells
            r2[0].text = str(i)
            r2[1].text = f"{x['giay']} s"
            r2[2].text = "có" if x["so_bo"] == "Cần người đọc" else "KHÔNG"
            r2[3].text = ", ".join(x["thieu"])[:220] or "—"

    if cham["thieu"] or cham["thua"]:
        d.add_paragraph("Dấu hiệu máy sàng:")
        for m in cham["thieu"]:
            d.add_paragraph(f"THIẾU dấu hiệu mong đợi: {m}", style="List Bullet")
        for m in cham["thua"]:
            d.add_paragraph(f"CÓ dấu hiệu KHÔNG được có: {m}", style="List Bullet")

    d.add_heading("4. Ảnh chụp màn hình từng bước", level=2)
    from docx.shared import Inches
    for a in kq["anh"]:
        d.add_paragraph(a)
        try:
            d.add_picture(str(thu_muc / a), width=Inches(6.3))
        except Exception as e:                                   # noqa: BLE001
            d.add_paragraph(f"(không nhúng được ảnh: {e})")

    d.add_heading("5. Nhật ký đầy đủ", level=2)
    d.add_paragraph("Nguyên văn, không cắt — một nhật ký bị cắt đẩy người đọc đi sửa "
                    "một lỗi không có thật.")
    for dong in kq["nhat_ky"].split("\n"):
        p = d.add_paragraph(dong)
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Menlo"
            run.font.size = Pt(7.5)

    d.save(str(dich))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--tc", nargs="*", help="chỉ chạy các mã TC này")
    ap.add_argument("--lai", action="store_true", help="chạy lại cả TC đã có kết quả")
    ap.add_argument("--lap", type=int, default=1,
                    help="chạy mỗi TC mấy lần (sheet Huong dan đòi 3–5 lần vì LLM không tất định)")
    a = ap.parse_args()

    ke = json.loads(KE_HOACH.read_text(encoding="utf-8"))
    chon = [t for t in ke if not a.tc or t["tc"] in a.tc]
    RA.mkdir(parents=True, exist_ok=True)

    # Dự án của TỪNG TC, không chỉ của TC liền trước. `sau: "TC003"` nói chuỗi này nối vào một
    # TC cụ thể; đi theo "cái vừa chạy" thì TC006 sẽ mở dự án của TC005 — một dự án nói về LoRa
    # — rồi hỏi nó về bộ chuyển LAN→USB, và đo một cuộc trao đổi chưa từng xảy ra.
    du_an: dict[str, Path] = {}
    du_an_truoc: Path | None = None
    ket: list[dict] = []
    for i, tc in enumerate(chon, 1):
        thu_muc = RA / tc["tc"]
        xong = thu_muc / "ket-qua.json"
        if xong.exists() and not a.lai:
            cu = json.loads(xong.read_text(encoding="utf-8"))
            if cu.get("du_an"):
                du_an[tc["tc"]] = du_an_truoc = Path(cu["du_an"])
            elif cu.get("noi_tiep") and du_an_truoc is not None:
                # Một TC nối tiếp KHÔNG tạo thư mục dự án riêng — nó mở lại dự án của TC gốc.
                # Không mang tiếp mốc ấy thì mắt xích sau nó ("TC006 nối TC003") mất điểm neo
                # ngay khi chạy lại một phần, và bộ chạy lặng lẽ đo trên một dự án khác.
                du_an[tc["tc"]] = du_an_truoc
            ket.append(cu)
            print(f"[{i}/{len(chon)}] {tc['tc']} — đã có, bỏ qua")
            continue

        if tc.get("noi_tiep"):
            goc = du_an.get(tc["sau"]) if tc.get("sau") else du_an_truoc
            if goc is None:
                print(f"     ⚠ {tc['tc']} nối tiếp {tc.get('sau') or 'TC trước'} nhưng TC ấy "
                      f"chưa chạy — chạy nó trước, nếu không đây là một phép đo mù")
            du_an_truoc = goc

        if tc.get("che_do") == "chan":
            kq = {"giay": 0.0, "rc": 0, "nhat_ky": tc.get("ly_do_chan", ""), "anh": [],
                  "du_an": None}
        else:
            if thu_muc.exists():
                shutil.rmtree(thu_muc)
            print(f"[{i}/{len(chon)}] {tc['tc']} — {tc['ten'][:56]} …", flush=True)
            # CHẠY LẶP — sheet "Huong dan": *"mỗi test nên chạy lặp 3–5 lần do kết quả không
            # tất định; ghi tỉ lệ đạt thay vì 1 lần"*. Đo 23/09/2026: TC003 qua sàng ở lần chạy
            # riêng và trượt ở lần chạy đủ bộ, cùng một bản build — một phán quyết rút từ một
            # lần chạy là một phán quyết không lặp lại được.
            lan: list[dict] = []
            for k in range(a.lap):
                tm = thu_muc if a.lap == 1 else thu_muc / f"lan-{k + 1}"
                try:
                    r1 = _chay_kich_ban(tc, tm, du_an_truoc)
                except subprocess.TimeoutExpired:
                    r1 = {"giay": -1.0, "rc": 124, "nhat_ky": "QUÁ HẠN — bộ lái không thoát.",
                          "anh": sorted(x.name for x in tm.glob("buoc-*.png")), "du_an": None}
                r1["cham"] = _cham(tc, r1["nhat_ky"])
                r1["thu_muc"] = str(tm)     # docx nhúng ảnh từ ĐÚNG thư mục của lần đó
                lan.append(r1)
                if a.lap > 1:
                    print(f"     lần {k + 1}/{a.lap}: {r1['giay']}s · {r1['cham']['so_bo']}",
                          flush=True)
            # Giữ lần ĐẦU làm bằng chứng chính để ảnh và nhật ký trong docx khớp nhau; các lần
            # sau chỉ góp vào tỉ lệ. Lấy "lần đẹp nhất" sẽ biến bộ đo thành bộ chọn kết quả.
            kq = lan[0]
            kq["lap"] = [{"giay": x["giay"], "so_bo": x["cham"]["so_bo"],
                          "thieu": x["cham"]["thieu"], "thua": x["cham"]["thua"]} for x in lan]
            kq["qua_sang"] = sum(1 for x in lan if x["cham"]["so_bo"] == "Cần người đọc")
            if kq.get("du_an"):
                du_an[tc["tc"]] = du_an_truoc = Path(kq["du_an"])
            elif tc.get("noi_tiep") and du_an_truoc is not None:
                du_an[tc["tc"]] = du_an_truoc          # chuỗi nối tiếp giữ nguyên dự án gốc

        # GOM TOÀN CẢNH ngay sau khi chạy, trước khi ca sau đụng vào gì. [DEV-217]
        # Bằng chứng của một ca nằm rải ở sáu nơi (nhật ký, stdout, ledger, store, session,
        # cấu hình); gom muộn thì một ca `noi_tiep` đã ghi đè lên trạng thái cần đọc.
        try:
            _tc = thu_toan_canh.thu(thu_muc, Path(kq["du_an"]) if kq.get("du_an") else None)
            thu_toan_canh.viet(_tc, thu_muc)
        except Exception as e:                                  # noqa: BLE001
            # Gom log hỏng KHÔNG được làm hỏng lượt đo.
            print(f"     ⚠ gom toàn cảnh {tc['tc']} hỏng: {e}", flush=True)

        cham = kq.get("cham") or _cham(tc, kq["nhat_ky"])
        if kq.get("lap") and len(kq["lap"]) > 1:
            tc["ti_le_qua_sang"] = f"{kq['qua_sang']}/{len(kq['lap'])}"
        tc.setdefault("trang_thai", cham["so_bo"])
        _docx(tc, kq, cham, Path(kq.get("thu_muc") or thu_muc), RA / f"{tc['tc']}.docx")

        d = dict(tc)
        d.update({"giay": kq["giay"], "rc": kq["rc"], "anh": kq["anh"],
                  "du_an": str(kq["du_an"]) if kq.get("du_an") else None,
                  "cham": cham, "lap": kq.get("lap"), "qua_sang": kq.get("qua_sang")})
        thu_muc.mkdir(parents=True, exist_ok=True)
        xong.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
        ket.append(d)
        print(f"     {kq['giay']}s · {len(kq['anh'])} ảnh · {cham['so_bo']}"
              + (f" · qua sàng {kq['qua_sang']}/{len(kq['lap'])}" if kq.get("lap")
                 and len(kq["lap"]) > 1 else "")
              + (f" · thiếu {cham['thieu']}" if cham["thieu"] else ""), flush=True)

    print(f"\n=== {len(ket)} TC đã chạy. docx trong {RA}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
