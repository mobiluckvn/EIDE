# -*- coding: utf-8 -*-
"""Xây EIDE_Danh_muc_Nang_luc.xlsx từ caps.py"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from collections import Counter
import caps

RED, GOLD, SLATE, ZEBRA = "B8121F", "F2B705", "2F4858", "F2F6FB"
thin = Side(style="thin", color="D0D5DD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
HDR_FILL = PatternFill("solid", fgColor=RED)
SUB_FILL = PatternFill("solid", fgColor=SLATE)
BODY = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = openpyxl.Workbook()

# ---------- Sheet 0 ----------
ws0 = wb.active; ws0.title = "0. Hướng dẫn"
ns_all = []
for c in caps.C:
    if c[1] not in ns_all: ns_all.append(c[1])
lines = [
 "EIDE — DANH MỤC NĂNG LỰC ĐẦY ĐỦ (CAPABILITY CATALOG)",
 "Mọi chức năng phần mềm là một năng lực có hợp đồng; tác tử hiểu lệnh → đối chiếu trạng thái → lập chuỗi → gọi năng lực theo chính sách tự chủ (EIDE-APD-08, DPS-09). 05/09/2026 · Vũ Trí Công (hỗ trợ: Claude)",
 "",
 "Cột",
 "Mã / Nhóm: không gian tên của năng lực (" + ", ".join(ns_all) + ").",
 "Lớp rủi ro R0 đọc · R1 ghi tri thức · R2 ghi mã · R3 phần cứng lab · R4 không hoàn tác/vật lý (APD §3).",
 "Mức: T1 AI làm trọn · T1* tự làm khi chính sách đủ bằng chứng, nếu không hỏi · T2 AI làm — người duyệt · T3 người làm.",
 "Grounding: điều tác tử phải đối chiếu trước khi gọi. Hỏi kỹ sư khi: điều kiện chuyển sang hỏi. Trạng thái M0: đã có trong hkw-core hay chưa.",
 "Sheet 2 thống kê bằng công thức theo nhóm/mức/lớp rủi ro/trạng thái. Sheet 3 xếp năng lực theo vòng đời kỹ nghệ (yêu cầu → kiến trúc → thiết kế → mã → kiểm thử → tài liệu).",
 "Bổ sung 05/09 (v1.2): nhóm tool (tác tử tự viết công cụ Python và chạy — năng lực gốc sinh năng lực khác); (v1.1): nhóm req (phân tích yêu cầu), view (bản đồ tri thức, RAG), discover (dò board, probe, tốc độ kết nối), arch (thiết kế kiến trúc), diagram (vẽ lược đồ bằng ngôn ngữ GEditor hỗ trợ: Mermaid, PlantUML, Graphviz/DOT, D2, WaveDrom, SVG), doc (viết tài liệu theo chuẩn bộ EAA/EIDE).",
]
for i, t in enumerate(lines, 1):
    c = ws0.cell(row=i, column=1, value=t); c.alignment = WRAP; c.font = BODY
ws0["A1"].font = Font(bold=True, size=14, color=RED, name="Arial")
ws0["A4"].font = Font(bold=True, name="Arial", size=10)
ws0.column_dimensions["A"].width = 150

# ---------- Sheet 1 ----------
ws1 = wb.create_sheet("1. Danh mục năng lực")
hdr = ["Mã", "Nhóm", "Năng lực", "Mô tả", "Tham số vào", "Kết quả ra", "Lớp rủi ro", "Mức",
       "Grounding / tiền điều kiện", "Hỏi kỹ sư khi", "UC / tham chiếu", "Mốc", "Trạng thái M0"]
ws1.append(hdr)
for c in ws1[1]:
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
ws1.row_dimensions[1].height = 23.25
for i, row in enumerate(caps.C, 2):
    ws1.append(list(row))
    for c in ws1[i]:
        c.font = BODY; c.alignment = WRAP; c.border = BORDER
        if i % 2 == 0: c.fill = PatternFill("solid", fgColor=ZEBRA)
n = len(caps.C) + 1
for col, w in zip("ABCDEFGHIJKLM", [11, 10, 26, 52, 24, 18, 9, 7, 26, 28, 16, 6, 14]):
    ws1.column_dimensions[col].width = w
ws1.freeze_panes = "D2"
t = Table(displayName="NangLucDayDu", ref=f"A1:M{n}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ws1.add_table(t)

# ---------- Sheet 2 ----------
ws2 = wb.create_sheet("2. Thống kê")
S = "'1. Danh mục năng lực'!"
def hdr_row(r, *vals):
    for j, v in enumerate(vals, 1):
        c = ws2.cell(row=r, column=j, value=v); c.font = HDR_FONT; c.fill = SUB_FILL
ws2["A1"] = "THỐNG KÊ (công thức)"; ws2["A1"].font = Font(bold=True, size=12, color=RED, name="Arial")
ws2["A2"] = "Tổng năng lực"; ws2["B2"] = f"=COUNTA({S}A2:A{n})"
ws2["A3"] = "Đã có trong M0 (chứa 'có')"; ws2["B3"] = f'=COUNTIF({S}M2:M{n},"*có*")'
hdr_row(5, "Theo mức", "Số", "Tỷ lệ")
for k, (lab, crit) in enumerate([("T1", '"T1"'), ("T1*", '"T1~*"'), ("T2", '"T2"'), ("T3", '"T3"')]):
    r = 6 + k
    ws2.cell(row=r, column=1, value=lab)
    ws2.cell(row=r, column=2, value=f"=COUNTIF({S}H2:H{n},{crit})")
    ws2.cell(row=r, column=3, value=f"=B{r}/$B$2").number_format = "0.0%"
hdr_row(11, "Theo lớp rủi ro", "Số")
for k, rc in enumerate(["R0", "R1", "R2", "R3", "R4"]):
    ws2.cell(row=12 + k, column=1, value=rc); ws2.cell(row=12 + k, column=2, value=f'=COUNTIF({S}G2:G{n},"{rc}*")')
hdr_row(18, "Theo nhóm", "Số", "T1", "T2", "T3")
r = 19
for g in sorted(ns_all):
    ws2.cell(row=r, column=1, value=g)
    ws2.cell(row=r, column=2, value=f'=COUNTIF({S}B2:B{n},"{g}")')
    ws2.cell(row=r, column=3, value=f'=COUNTIFS({S}B2:B{n},"{g}",{S}H2:H{n},"T1~*")+COUNTIFS({S}B2:B{n},"{g}",{S}H2:H{n},"T1")')
    ws2.cell(row=r, column=4, value=f'=COUNTIFS({S}B2:B{n},"{g}",{S}H2:H{n},"T2")')
    ws2.cell(row=r, column=5, value=f'=COUNTIFS({S}B2:B{n},"{g}",{S}H2:H{n},"T3")')
    r += 1
r += 1
hdr_row(r, "Theo mốc", "Số"); r += 1
for m in ["M0", "M1", "M2", "M3", "M4", "M5"]:
    ws2.cell(row=r, column=1, value=m); ws2.cell(row=r, column=2, value=f'=COUNTIF({S}L2:L{n},"{m}")'); r += 1
ws2.column_dimensions["A"].width = 28; ws2.column_dimensions["B"].width = 10

# ---------- Sheet 3: vòng đời kỹ nghệ ----------
ws3 = wb.create_sheet("3. Theo vòng đời")
LIFE = [
 ("1. Thu thập & phân tích yêu cầu", ["req", "chat", "memory", "search"], "Kỹ sư ra lệnh; tác tử thu thập, phân loại, đối chiếu khả thi với hộ chiếu; câu hỏi gộp cho ô trống; sinh URD/SRS."),
 ("2. Tri thức phần cứng", ["archive", "extract", "passport", "kg", "board", "view"], "Khai phá zip/rar, trích xuất PDF/ảnh/BOM/netlist → fact có nguồn → hộ chiếu chip/board và đồ thị tri thức; hiển thị bản đồ tri thức, nguồn gốc, hỏi–đáp RAG có trích dẫn."),
 ("3. Thiết kế kiến trúc & chi tiết", ["arch", "diagram", "plan"], "Chọn kiểu kiến trúc, phân rã module, gán tài nguyên phần cứng, ngân sách RAM/thời gian, ADR; vẽ sơ đồ khối/kiến trúc/tuần tự/trạng thái; lập kế hoạch."),
 ("4. Môi trường & hiện thực", ["env", "code", "project"], "Kiểm tra/cài công cụ (hỏi khi R2+), sinh mã theo module, tích hợp, merge có hoàn tác."),
 ("5. Mô phỏng, nạp, gỡ lỗi, đo", ["discover", "sim", "target", "debug", "measure", "bench"], "Dò board/probe/cổng, ID chip, tốc độ kết nối, quét bus; dựng mô phỏng, nạp/verify qua ISA adapter, thí nghiệm theo chính sách, đo đạc, benchmark."),
 ("6. Tài liệu, báo cáo, phát hành", ["doc", "report", "registry"], "Sinh bộ tài liệu chuẩn EAA/EIDE có mục Nguồn, chèn lược đồ, báo cáo kiểm thử/tiến độ, đóng gói .hkp."),
 ("Xuyên suốt: chính sách tự chủ", ["policy"], "APPROVE/ASK/REJECT theo cổng, cửa sổ hoàn tác, leo thang, dừng khẩn, học ngưỡng."),
 ("Xuyên suốt: tự tạo công cụ (self-tooling)", ["tool"], "Khi thiếu năng lực phù hợp, tác tử tự đặc tả, viết công cụ Python, test trong sandbox, chạy theo chính sách rủi ro suy ra từ hiệu ứng, đăng ký thành năng lực tạm và thăng cấp sau khi chứng minh — năng lực gốc sinh ra năng lực khác."),
]
ws3.append(["Giai đoạn", "Nhóm năng lực", "Số năng lực", "Mô tả", "Mã năng lực"])
for c in ws3[1]:
    c.font = HDR_FONT; c.fill = HDR_FILL; c.border = BORDER; c.alignment = Alignment(wrap_text=True, vertical="center")
cnt = Counter(c[1] for c in caps.C)
for i, (stage, groups, desc) in enumerate(LIFE, 2):
    ids = ", ".join(c[2] for c in caps.C if c[1] in groups)
    ws3.append([stage, ", ".join(groups), sum(cnt[g] for g in groups), desc, ids])
    for c in ws3[i]:
        c.font = BODY; c.alignment = WRAP; c.border = BORDER
        if i % 2 == 0: c.fill = PatternFill("solid", fgColor=ZEBRA)
for col, w in zip("ABCDE", [30, 26, 12, 60, 90]):
    ws3.column_dimensions[col].width = w
ws3.freeze_panes = "B2"

wb.save("EIDE_Danh_muc_Nang_luc.xlsx")
print("caps:", len(caps.C), dict(cnt))
