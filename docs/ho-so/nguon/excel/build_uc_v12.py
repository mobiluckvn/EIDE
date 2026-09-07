# -*- coding: utf-8 -*-
"""Use case v1.2: nạp workbook hiện có, thêm 11 UC mới vào sheet 2/3/4, thêm sheet 10 (sequence Mermaid) và 11 (bước ↔ năng lực)."""
import openpyxl, re, json
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
import uc_data as D
import uc_data_v12  # đăng ký UC mới vào D.UCS/FLOWS/ALTS
caps = json.load(open("caps.json", encoding="utf-8"))
capnames = {c["name"] for c in caps}
RED, SLATE, ZEBRA, LIGHT = "B8121F", "2F4858", "F2F6FB", "FFF4D6"
thin = Side(style="thin", color="D0D5DD"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = Font(bold=True, color="FFFFFF", name="Arial", size=10); BODY = Font(name="Arial", size=10); WRAP = Alignment(wrap_text=True, vertical="top")

wb = openpyxl.load_workbook("EIDE_Use_Case_Chi_Tiet.xlsx")
ws0 = wb.worksheets[0]
ws0.insert_rows(3); ws0.cell(row=3, column=1, value="v1.2 (05/09/2026): thêm 11 use case nhóm I (kỹ nghệ), J (tri thức & dò tìm), K (tự tạo công cụ) — 56 UC; sheet 10 sơ đồ tuần tự Mermaid cho mọi UC; sheet 11 ánh xạ bước ↔ năng lực ↔ màn hình; cổng theo chính sách (APD-08/POL-17).").font = Font(name="Arial", size=10, bold=True, color=RED)
NEW = [u for u in D.UCS if u[0][3] in "IJK"]
new_ids = {u[0] for u in NEW}

def style_row(ws, r, zebra):
    for c in ws[r]:
        c.font = BODY; c.alignment = WRAP; c.border = BORDER
        if zebra: c.fill = PatternFill("solid", fgColor=ZEBRA)

# ---- sheet 2: Danh mục UC (cột như build_xlsx: 16 trường + số bước? kiểm tra header)
ws2 = wb["2. Danh mục UC"]; hdr2 = [c.value for c in ws2[1]]
for u in NEW:
    row = list(u) + [len(D.FLOWS[u[0]]), len(D.ALTS[u[0]])]
    row = row[:len(hdr2)] + [None] * (len(hdr2) - len(row))
    ws2.append(row); style_row(ws2, ws2.max_row, ws2.max_row % 2 == 0)
# ---- sheet 3: Luồng chính
ws3 = wb["3. Luồng chính"]
for u in NEW:
    for i, (lane, act, inp, out, note) in enumerate(D.FLOWS[u[0]], 1):
        ws3.append([u[0], u[1], i, lane, act, inp, out, note]); style_row(ws3, ws3.max_row, ws3.max_row % 2 == 0)
        if lane == "N" or str(note).startswith("G"):
            for c in ws3[ws3.max_row]: c.fill = PatternFill("solid", fgColor=LIGHT)
# ---- sheet 4: Luồng thay thế
ws4 = wb["4. Luồng thay thế"]
for u in NEW:
    for a, cond, h in D.ALTS[u[0]]:
        ws4.append([u[0], u[1], a, cond, h]); style_row(ws4, ws4.max_row, ws4.max_row % 2 == 0)
# cập nhật ref bảng
for ws, tname in ((ws2, "DanhMucUC"), (ws3, "LuongChinh"), (ws4, "LuongThayThe")):
    if tname in ws.tables:
        t = ws.tables[tname]; t.ref = f"A1:{t.ref.split(':')[1].rstrip('0123456789')}{ws.max_row}"

# ---- sheet 10: Sequence Mermaid
LANE = {"N": "KyS as Kỹ sư", "T": "TT as Tác tử", "K": "KP as Knowledge Plane", "H": "HW as Công cụ/Phần cứng", "N/T": "KyS as Kỹ sư", "T/K": "TT as Tác tử", "T/H": "TT as Tác tử", "N/H": "KyS as Kỹ sư", "T/N": "TT as Tác tử", "K/T": "KP as Knowledge Plane"}
ALIAS = {"N": "KyS", "T": "TT", "K": "KP", "H": "HW"}
def mermaid(uid):
    steps = D.FLOWS[uid]; lines = ["sequenceDiagram", "  autonumber"]
    used = []
    for lane, *_ in steps:
        for l in lane.split("/"):
            if l in ALIAS and ALIAS[l] not in used: used.append(ALIAS[l])
    for a in ["KyS", "TT", "KP", "HW"]:
        if a in used: lines.append("  participant " + {"KyS": "KyS as Kỹ sư", "TT": "TT as Tác tử", "KP": "KP as Knowledge Plane", "HW": "HW as Công cụ/Phần cứng"}[a])
    for idx, (lane, act, inp, out, note) in enumerate(steps):
        src = ALIAS[lane.split("/")[0]]
        if "/" in lane: dst = ALIAS[lane.split("/")[-1]]
        else:
            nxt = steps[idx + 1][0].split("/")[0] if idx + 1 < len(steps) else None
            dst = ALIAS.get(nxt, None) if nxt else None
        if not dst or dst == src: dst = "KP" if src != "KP" else "TT"
        txt = re.sub(r"[;:]", ",", act)[:90]
        gate = f" [{note}]" if note and (str(note).startswith("G") or note in ("Lệnh", "T1*", "T2", "T3", "D3")) else ""
        lines.append(f"  {src}->>{dst}: {txt}{gate}")
        if out: lines.append(f"  Note over {dst}: ra: {out[:60]}")
    return "\n".join(lines)
ws10 = wb.create_sheet("10. Sequence (Mermaid)")
ws10.append(["Mã UC", "Tên", "Mermaid sequenceDiagram (dán vào GEditor/DiagramView)"])
for c in ws10[1]: c.font = HF; c.fill = PatternFill("solid", fgColor=RED); c.border = BORDER
for u in D.UCS:
    ws10.append([u[0], u[1], mermaid(u[0])]); style_row(ws10, ws10.max_row, ws10.max_row % 2 == 0)
    ws10.cell(row=ws10.max_row, column=3).font = Font(name="Consolas", size=9)
for col, w in zip("ABC", [10, 40, 120]): ws10.column_dimensions[col].width = w
ws10.freeze_panes = "C2"

# ---- sheet 11: Bước ↔ năng lực ↔ màn hình
KEYMAP = [("dự án", "project.create"), ("mở nén|zip|giải nén", "archive.unpack"), ("phân loại", "ingest.classify"), ("tìm ứng viên|registry → ", "search.web"), ("tải", "search.fetch"), ("svd", "extract.svd"), ("pdf", "extract.pdf_register_map"), ("ảnh", "extract.image_schematic"), ("kicad|netlist", "extract.kicad_netlist"), ("bom", "extract.bom"),
          ("duyệt fact|g-fact", "kg.review_facts"), ("mâu thuẫn", "kg.resolve_conflict"), ("kiểm định", "passport.verify_on_board"), ("hộ chiếu mạch|boardpassport", "board.build_passport"), ("xung đột", "board.check_pins"), ("doctor|cài", "env.install"), ("kế hoạch|g1", "plan.create"), ("sinh mã|coder", "code.generate_module"), ("constant-guard", "code.constant_guard"), ("build|biên dịch", "code.build"), ("review|g3", "code.review"), ("merge", "code.merge"),
          ("mô phỏng|sil|renode", "sim.run"), ("nạp|flash", "target.flash"), ("serial|expect", "target.serial"), ("probe", "target.probe_read"), ("g4|xác nhận vật lý|quan sát", "target.observe"), ("log", "debug.log_stats"), ("giả thuyết", "debug.hypothesize"), ("thí nghiệm", "debug.experiment"), ("đóng gói|.hkp", "registry.pack"), ("phát hành|g5", "registry.publish"), ("pull", "registry.pull"), ("benchmark", "bench.run"),
          ("báo cáo", "report.progress"), ("quyền|g-ops", "policy.permit"), ("lệnh", "chat.parse_intent"), ("hỏi", "chat.clarify"), ("đủ thông tin", "plan.sufficiency"), ("từ chối|không biết", "chat.decline"), ("rollback|known-good", "project.rollback"), ("tiếp tục phiên|progress", "memory.summarize_session")]
SCREEN = {u[0]: u[15] for u in D.UCS}
ws11 = wb.create_sheet("11. Bước ↔ năng lực")
ws11.append(["Mã UC", "Bước", "Làn", "Hành động", "Năng lực gọi (khai báo hoặc suy ra)", "Màn hình", "Cổng / mức"])
for c in ws11[1]: c.font = HF; c.fill = PatternFill("solid", fgColor=RED); c.border = BORDER
n_explicit = 0
for u in D.UCS:
    for i, (lane, act, inp, out, note) in enumerate(D.FLOWS[u[0]], 1):
        found = sorted({m for m in re.findall(r"\b[a-z]+\.[a-z_]+\b", act + " " + str(note)) if m in capnames})
        if found: n_explicit += 1
        else:
            low = act.lower()
            for pat, capn in KEYMAP:
                if re.search(pat, low): found = [capn + " (suy ra)"]; break
        ws11.append([u[0], i, lane, act, ", ".join(found) if found else "— (bước của người / không gọi năng lực)", SCREEN[u[0]], note]); style_row(ws11, ws11.max_row, ws11.max_row % 2 == 0)
for col, w in zip("ABCDEFG", [10, 6, 6, 70, 34, 20, 16]): ws11.column_dimensions[col].width = w
ws11.freeze_panes = "D2"
wb.save("EIDE_Use_Case_Chi_Tiet_v1.2.xlsx")
print("UC", len(D.UCS), "new", len(NEW), "steps", ws11.max_row - 1, "explicit", n_explicit)
