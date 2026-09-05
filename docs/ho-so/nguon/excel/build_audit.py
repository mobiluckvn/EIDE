# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
import caps, uc_data
from audit_data import LEVELS, AREAS, DOCS, NS_ROWS, SCREENS, UC_GAPS, NEW_UCS, NEW_DOCS

RED, SLATE, ZEBRA, GREEN, AMBER, ROSE = "B8121F", "2F4858", "F2F6FB", "E6F4EA", "FFF4D6", "FDECEE"
thin = Side(style="thin", color="D0D5DD"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = Font(bold=True, color="FFFFFF", name="Arial", size=10); HFILL = PatternFill("solid", fgColor=RED)
BODY = Font(name="Arial", size=10); WRAP = Alignment(wrap_text=True, vertical="top")
wb = openpyxl.Workbook()

def sheet(title, header, rows, widths, freeze="B2", table=None):
    ws = wb.create_sheet(title); ws.append(header)
    for c in ws[1]: c.font = HF; c.fill = HFILL; c.border = BORDER; c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, 2):
        ws.append(list(r))
        for c in ws[i]:
            c.font = BODY; c.alignment = WRAP; c.border = BORDER
            if i % 2 == 0: c.fill = PatternFill("solid", fgColor=ZEBRA)
    for col, w in zip("ABCDEFGHIJKLMNOP", widths): ws.column_dimensions[col].width = w
    ws.freeze_panes = freeze
    if table:
        t = Table(displayName=table, ref=f"A1:{chr(64 + len(header))}{len(rows) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False); ws.add_table(t)
    return ws

def status_colors(ws, col, n):
    rng = f"{col}2:{col}{n + 1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Đủ"'], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Một phần"'], fill=PatternFill("solid", fgColor=AMBER)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Thiếu"'], fill=PatternFill("solid", fgColor=ROSE)))

# ---- Sheet 0 ----
ws0 = wb.active; ws0.title = "0. Hướng dẫn"
lines = [
 ("EIDE — RÀ SOÁT ĐỘ ĐẦY ĐỦ BỘ HỒ SƠ THIẾT KẾ ĐỂ ĐƯA VÀO PHÁT TRIỂN", True),
 ("Đề án tốt nghiệp Thạc sĩ ngành Kỹ thuật Điện tử — Học viện Công nghệ Bưu chính Viễn thông (PTIT) · Học viên: Vũ Trí Công · Người hướng dẫn: TS. Nguyễn Trung Hiếu · 05/09/2026 (hỗ trợ: Claude)", False),
 ("Câu hỏi: bộ hồ sơ v1.1 (10 tài liệu + 2 Excel + 17 mockup + mã M0) đã chi tiết đến MỨC 3 (lập trình được) cho toàn bộ hệ thống chưa? Nếu chưa, thiếu gì và cần tài liệu nào.", False),
 ("", False),
 ("Định nghĩa ba mức chi tiết", True),
]
for l in LEVELS: lines.append((f"{l[0]} — {l[1]}: {l[2]}", False))
lines += [("", False), ("Các sheet", True),
 ("1. Lĩnh vực: 40 lĩnh vực thiết kế cần có để phát triển, mức hiện tại, kết luận Đủ / Một phần / Thiếu, khoảng trống cụ thể, tài liệu đề xuất, ưu tiên, mốc cần trước, ước lượng.", False),
 ("2. Tài liệu hiện có: 14 hiện vật (10 docx, 2 xlsx, mockup, mã) — mức đạt, còn thiếu gì, hướng xử lý.", False),
 ("3. Nhóm năng lực: 26 nhóm (228 năng lực) — mức chi tiết hiện có, thiếu gì ở mức 3, tập CDS-12 đích, mốc.", False),
 ("4. UI/UX: 22 màn hình (17 có mockup + 5 mới) — cập nhật v1.1 cần làm, đặc tả còn thiếu, năng lực chính.", False),
 ("5. Use case: 45 UC hiện có + 10 UC mới (nhóm I, J) — có gì, thiếu gì.", False),
 ("6. Kế hoạch tài liệu bổ sung: 19 tài liệu/phụ lục mới (CXD-10 … CON-28) với mục lục, đầu vào, mốc, thứ tự làm, ước lượng ngày, trạng thái.", False),
 ("7. Thống kê: công thức tổng hợp — tỷ lệ lĩnh vực đủ mức 3, số ngày ước lượng theo mốc, tiến độ tài liệu bổ sung.", False),
 ("", False), ("Kết luận ngắn", True),
 ("Đã đủ mức 3: định vị/nguyên tắc, yêu cầu người dùng, kiến trúc hệ thống, luồng nghiệp vụ, kiến trúc tri thức, chính sách hội thoại, lõi dữ liệu (store/KG/gateway — có mã M0), hợp đồng năng lực và Router.", False),
 ("Mới ở mức 2 (phải xuống mức 3 trước khi code): đặc tả chi tiết 228 năng lực (nhất là extract, code, diagram, view, discover, arch), chính sách tự chủ máy đọc được, API đầy đủ, từ điển dữ liệu/DDL, ISA/toolchain/adapter, benchmark, gói registry, kiểm thử fixture.", False),
 ("Thiếu hẳn (mức 1): kiến trúc ngữ cảnh, kiến trúc bộ nhớ tác tử, prompt & vai trò, đặc tả UI/UX (mockup chưa theo v1.1), thiết kế mô phỏng, giao diện plugin GEditor, triển khai/cài đặt.", False),
 ("Đề xuất: làm chậm và đủ — 19 tài liệu/phụ lục bổ sung (~77 ngày công ước lượng, có AI hỗ trợ giảm ~60%), theo thứ tự ở sheet 6; nhóm phải xong TRƯỚC khi bắt đầu mã M1: CXD-10, MEM-11, UXD-13, DDD-14, CDS-12 (tập 2, 3, 5, 6), API-15, PRS-16, POL-17, UCD-24, PLN-27.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws0.cell(row=i, column=1, value=t); c.alignment = WRAP; c.font = Font(name="Arial", size=10, bold=b, color=RED if (b and i == 1) else "000000")
ws0["A1"].font = Font(bold=True, size=14, color=RED, name="Arial"); ws0.column_dimensions["A"].width = 160

# ---- Sheet 1 ----
h1 = ["Mã", "Lĩnh vực thiết kế", "Cần gì để đạt mức 3", "Tài liệu hiện có", "Mức hiện tại (1–3)", "Kết luận", "Khoảng trống cụ thể", "Tài liệu đề xuất", "Ưu tiên", "Mốc cần trước", "Ước lượng (ngày)"]
ws1 = sheet("1. Lĩnh vực", h1, AREAS, [6, 34, 38, 30, 9, 11, 60, 34, 8, 9, 10], "C2", "LinhVuc"); status_colors(ws1, "F", len(AREAS))

# ---- Sheet 2 ----
h2 = ["Hiện vật", "Tên / phiên bản", "Mức đạt", "Kết luận", "Nội dung chính đã có", "Còn thiếu để phát triển", "Hướng xử lý"]
ws2 = sheet("2. Tài liệu hiện có", h2, DOCS, [26, 30, 16, 20, 50, 55, 30], "B2", "TaiLieuHienCo")

# ---- Sheet 3 ----
rows3 = NS_ROWS(caps.C)
h3 = ["Nhóm", "Số năng lực", "Đã có mã M0", "Mức hiện tại", "Kết luận", "Thiếu ở mức 3", "Tài liệu đích", "Mốc"]
ws3 = sheet("3. Nhóm năng lực", h3, rows3, [12, 10, 10, 10, 11, 70, 30, 9], "B2", "NhomNangLuc"); status_colors(ws3, "E", len(rows3))
r = len(rows3) + 3
ws3.cell(row=r, column=1, value="Tổng").font = Font(bold=True, name="Arial", size=10)
ws3.cell(row=r, column=2, value=f"=SUM(B2:B{len(rows3) + 1})"); ws3.cell(row=r, column=3, value=f"=SUM(C2:C{len(rows3) + 1})")
ws3.cell(row=r + 1, column=1, value="Năng lực trong nhóm đã đủ mức 3").font = Font(bold=True, name="Arial", size=10)
ws3.cell(row=r + 1, column=2, value=f'=SUMIF(E2:E{len(rows3) + 1},"Đủ",B2:B{len(rows3) + 1})')
ws3.cell(row=r + 2, column=1, value="Tỷ lệ năng lực đã đủ mức 3").font = Font(bold=True, name="Arial", size=10)
ws3.cell(row=r + 2, column=2, value=f"=B{r + 1}/B{r}").number_format = "0.0%"

# ---- Sheet 4 ----
h4 = ["Màn hình", "Có mockup", "Mức hiện tại", "Cập nhật v1.1 cần làm", "Đặc tả còn thiếu", "Năng lực chính"]
ws4 = sheet("4. UI-UX", h4, SCREENS, [26, 10, 10, 60, 40, 36], "B2", "ManHinh")
r = len(SCREENS) + 3
ws4.cell(row=r, column=1, value="Màn hình chưa có mockup").font = Font(bold=True, name="Arial", size=10); ws4.cell(row=r, column=2, value=f'=COUNTIF(B2:B{len(SCREENS) + 1},"Chưa")')
ws4.cell(row=r + 1, column=1, value="Màn hình cần cập nhật theo v1.1").font = Font(bold=True, name="Arial", size=10); ws4.cell(row=r + 1, column=2, value=f'=COUNTA(D2:D{len(SCREENS) + 1})-COUNTIF(D2:D{len(SCREENS) + 1},"—")')

# ---- Sheet 5 ----
rows5 = []
for u in uc_data.UCS:
    g = u[0][3]; have, gap = UC_GAPS[g]
    rows5.append((u[0], u[1], u[2], "Có", have, gap, "Cập nhật v1.2"))
for u in NEW_UCS:
    rows5.append((u[0], u[1], u[2], "Chưa", "—", "Viết mới toàn bộ (luồng, thay thế, cổng, mức tự chủ, năng lực, màn hình, sequence)", "Viết mới; năng lực: " + u[3]))
h5 = ["Mã UC", "Use case", "Nhóm", "Có trong Excel", "Đã có", "Còn thiếu", "Hành động"]
ws5 = sheet("5. Use case", h5, rows5, [9, 44, 24, 10, 36, 50, 34], "C2", "UseCase")
r = len(rows5) + 3
ws5.cell(row=r, column=1, value="UC hiện có").font = Font(bold=True, name="Arial", size=10); ws5.cell(row=r, column=2, value=f'=COUNTIF(D2:D{len(rows5) + 1},"Có")')
ws5.cell(row=r + 1, column=1, value="UC cần viết mới").font = Font(bold=True, name="Arial", size=10); ws5.cell(row=r + 1, column=2, value=f'=COUNTIF(D2:D{len(rows5) + 1},"Chưa")')

# ---- Sheet 6 ----
rows6 = [list(d) + ["Chưa bắt đầu"] for d in NEW_DOCS]
h6 = ["Mã", "Tên tài liệu", "Mục đích", "Nội dung chính (mục lục)", "Đầu vào", "Mức đích", "Mốc cần trước", "Thứ tự làm", "Ước lượng (ngày)", "Dạng", "Trạng thái"]
ws6 = sheet("6. Tài liệu bổ sung", h6, rows6, [13, 34, 40, 80, 28, 8, 10, 8, 9, 20, 13], "C2", "TaiLieuBoSung")
n6 = len(rows6) + 1
ws6.cell(row=n6 + 2, column=1, value="Tổng ngày ước lượng").font = Font(bold=True, name="Arial", size=10); ws6.cell(row=n6 + 2, column=2, value=f"=SUM(I2:I{n6})")
ws6.cell(row=n6 + 3, column=1, value="Có AI hỗ trợ (×0,4)").font = Font(bold=True, name="Arial", size=10); ws6.cell(row=n6 + 3, column=2, value=f"=ROUND(B{n6 + 2}*0.4,0)")
ws6.cell(row=n6 + 4, column=1, value="Đã hoàn thành").font = Font(bold=True, name="Arial", size=10); ws6.cell(row=n6 + 4, column=2, value=f'=COUNTIF(K2:K{n6},"Hoàn thành")')
from openpyxl.worksheet.datavalidation import DataValidation
dv = DataValidation(type="list", formula1='"Chưa bắt đầu,Đang làm,Hoàn thành"', allow_blank=True); ws6.add_data_validation(dv); dv.add(f"K2:K{n6}")

# ---- Sheet 7 ----
ws7 = wb.create_sheet("7. Thống kê"); S1 = "'1. Lĩnh vực'!"; S6 = "'6. Tài liệu bổ sung'!"
nA = len(AREAS) + 1
ws7["A1"] = "THỐNG KÊ RÀ SOÁT (công thức)"; ws7["A1"].font = Font(bold=True, size=12, color=RED, name="Arial")
def hdr(r, *v):
    for j, x in enumerate(v, 1):
        c = ws7.cell(row=r, column=j, value=x); c.font = HF; c.fill = PatternFill("solid", fgColor=SLATE)
hdr(3, "Lĩnh vực theo kết luận", "Số", "Tỷ lệ")
for k, lab in enumerate(["Đủ", "Một phần", "Thiếu"]):
    ws7.cell(row=4 + k, column=1, value=lab); ws7.cell(row=4 + k, column=2, value=f'=COUNTIF({S1}F2:F{nA},"{lab}")'); ws7.cell(row=4 + k, column=3, value=f"=B{4 + k}/COUNTA({S1}A2:A{nA})").number_format = "0.0%"
ws7.cell(row=7, column=1, value="Tổng lĩnh vực"); ws7.cell(row=7, column=2, value=f"=COUNTA({S1}A2:A{nA})")
hdr(9, "Lĩnh vực theo mức hiện tại", "Số")
for k, lv in enumerate([1, 2, 3]):
    ws7.cell(row=10 + k, column=1, value=f"Mức {lv}"); ws7.cell(row=10 + k, column=2, value=f"=COUNTIF({S1}E2:E{nA},{lv})")
hdr(14, "Ngày ước lượng theo mốc cần trước (lĩnh vực)", "Ngày")
for k, m in enumerate(["M0", "M1", "M1–M2", "M2", "M3", "M4"]):
    ws7.cell(row=15 + k, column=1, value=m); ws7.cell(row=15 + k, column=2, value=f'=SUMIF({S1}J2:J{nA},"{m}",{S1}K2:K{nA})')
ws7.cell(row=21, column=1, value="Tổng ngày (lĩnh vực)"); ws7.cell(row=21, column=2, value=f"=SUM({S1}K2:K{nA})")
hdr(23, "Tài liệu bổ sung theo ưu tiên M/S/C (lĩnh vực)", "Số")
for k, p in enumerate(["M", "S", "C"]):
    ws7.cell(row=24 + k, column=1, value=p); ws7.cell(row=24 + k, column=2, value=f'=COUNTIF({S1}I2:I{nA},"{p}")')
hdr(28, "Tài liệu bổ sung (sheet 6)", "Giá trị")
ws7.cell(row=29, column=1, value="Số tài liệu"); ws7.cell(row=29, column=2, value=f"=COUNTA({S6}A2:A{n6})")
ws7.cell(row=30, column=1, value="Tổng ngày ước lượng"); ws7.cell(row=30, column=2, value=f"=SUM({S6}I2:I{n6})")
ws7.cell(row=31, column=1, value="Cần xong trước M1"); ws7.cell(row=31, column=2, value=f'=COUNTIF({S6}G2:G{n6},"M1*")')
ws7.cell(row=32, column=1, value="Đã hoàn thành"); ws7.cell(row=32, column=2, value=f'=COUNTIF({S6}K2:K{n6},"Hoàn thành")')
ws7.cell(row=33, column=1, value="Tiến độ"); ws7.cell(row=33, column=2, value="=B32/B29").number_format = "0.0%"
ws7.column_dimensions["A"].width = 46; ws7.column_dimensions["B"].width = 12; ws7.column_dimensions["C"].width = 10
wb.save("EIDE_Ra_soat_Du_dieu_kien_Phat_trien.xlsx"); print("areas", len(AREAS), "docs", len(DOCS), "ns", len(rows3), "screens", len(SCREENS), "uc", len(rows5), "newdocs", len(rows6))
