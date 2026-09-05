# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import uc_data as D

RED, GOLD, GREY, LIGHT = "B8121F", "F2B705", "F2F2F2", "FFF4E6"
FONT = "Arial"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=RED)
body_font = Font(name=FONT, size=10)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

def sheet(title, headers, rows, widths, table_name, freeze="A2", zebra=True):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = body_font; c.alignment = wrap; c.border = border
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 30
    if ws.max_row > 1:
        ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        t = Table(displayName=table_name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=zebra)
        ws.add_table(t)
    return ws

# ---------------- 0. Hướng dẫn
ws0 = wb.active; ws0.title = "0. Hướng dẫn"
lines = [
    ("EIDE — DANH MỤC USE CASE CHI TIẾT", Font(name=FONT, bold=True, size=14, color=RED)),
    ("Embedded IDE trên nền GEditor · Knowledge Plane cục bộ · Bộ hồ sơ HKW/EIDE v1.0 · 05/09/2026 · Người lập: Vũ Trí Công (hỗ trợ: Claude)", Font(name=FONT, italic=True, size=10)),
    ("", None),
    ("Cách đọc bộ sổ tính này", Font(name=FONT, bold=True, size=11)),
    ("1. Tác nhân & cổng — ai tham gia, cổng người nào tồn tại và đứng ở đâu.", body_font),
    ("2. Danh mục UC — mỗi dòng một use case: mục tiêu, kích hoạt, tiền/hậu điều kiện, cổng, tri thức sinh ra, ưu tiên (M/S/C), mốc (M0–M5), truy vết FR/UR, màn hình UI.", body_font),
    ("3. Luồng chính — từng bước của mỗi UC: làn (N = người, T = tác tử, K = Knowledge Plane, H = công cụ/phần cứng), hành động, đầu vào → đầu ra, cổng/ghi chú.", body_font),
    ("4. Luồng thay thế — điều kiện ngoại lệ và cách xử lý cho mỗi UC.", body_font),
    ("5. Kịch bản mẫu — tình huống anh nêu (một board mới, zip hỗn hợp, thiếu thông tin → tìm mạng có phê duyệt, dựng mô phỏng) ánh xạ thành chuỗi UC theo thời gian.", body_font),
    ("6. Truy vết — ma trận UC ↔ nhóm quy trình ↔ FR ↔ UR ↔ màn hình, kèm thống kê bằng công thức.", body_font),
    ("", None),
    ("Quy ước", Font(name=FONT, bold=True, size=11)),
    ("Mã UC: UC-<nhóm><số>; nhóm A dự án & môi trường, B nhận tri thức, C mạch & tài nguyên, D kế hoạch & sinh mã, E mô phỏng, F xác minh & gỡ lỗi, G chia sẻ & đánh giá, H tương tác tác tử.", body_font),
    ("Ưu tiên MoSCoW: M bắt buộc, S nên có, C có thể có. Mốc M0–M5 theo lộ trình HKW-PDA-00 §12.", body_font),
    ("Ô nền vàng nhạt trong sheet 3 = bước có cổng người (không thể bỏ qua bằng phần mềm).", body_font),
    ("FR/UR tham chiếu HKW-SRS-02 và HKW-URD-01 (một số UR tham chiếu EAA-URD-06 khi ghi chú 'EAA').", body_font),
    ("Sheet 2, 3, 4 là bảng Excel (Table) — dùng bộ lọc trên tiêu đề để lọc theo nhóm, tác nhân, cổng, mốc.", body_font),
    ("Nguồn: HKW-PDA-00, URD-01, SRS-02, SAD-03, SDD-04, STP-05, BPD-06, KAD-07 (05/09/2026); mã nguồn hkw-core M0.", Font(name=FONT, italic=True, size=9, color="595959")),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws0.cell(row=i, column=1, value=t)
    if f: c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws0.column_dimensions["A"].width = 150

# ---------------- 1. Tác nhân & cổng
ws1 = wb.create_sheet("1. Tác nhân & cổng")
ws1.append(["TÁC NHÂN", "", ""])
ws1["A1"].font = Font(name=FONT, bold=True, size=12, color=RED)
ws1.append(["Tác nhân", "Loại", "Vai trò"])
for c in ws1[2]:
    c.font = hdr_font; c.fill = hdr_fill; c.border = border
for a in D.ACTORS:
    ws1.append(list(a))
r = ws1.max_row + 2
ws1.cell(row=r, column=1, value="CỔNG NGƯỜI (GATE)").font = Font(name=FONT, bold=True, size=12, color=RED)
ws1.append(["Mã", "Tên", "Đứng ở đâu"])
for c in ws1[ws1.max_row]:
    c.font = hdr_font; c.fill = hdr_fill; c.border = border
for g in D.GATES:
    ws1.append(list(g))
for row in ws1.iter_rows(min_row=3, max_row=ws1.max_row):
    for c in row:
        if c.value is not None and c.font != hdr_font:
            c.font = body_font; c.alignment = wrap; c.border = border
ws1.column_dimensions["A"].width = 26; ws1.column_dimensions["B"].width = 22; ws1.column_dimensions["C"].width = 100

# ---------------- 2. Danh mục UC
H2 = ["Mã UC", "Tên use case", "Nhóm / quy trình", "Tác nhân chính", "Tác nhân phụ", "Mục tiêu", "Kích hoạt", "Tiền điều kiện", "Hậu điều kiện", "Cổng người", "Tri thức sinh ra", "Ưu tiên", "Mốc", "FR liên quan", "UR liên quan", "Màn hình UI", "Số bước", "Số luồng thay thế"]
rows2 = []
for u in D.UCS:
    uid = u[0]
    rows2.append(list(u) + [len(D.FLOWS[uid]), len(D.ALTS[uid])])
sheet("2. Danh mục UC", H2, rows2, [10, 34, 20, 16, 22, 40, 30, 26, 36, 14, 24, 8, 7, 22, 20, 20, 8, 10], "DanhMucUC", freeze="C2")

# ---------------- 3. Luồng chính
H3 = ["Mã UC", "Tên use case", "Bước", "Làn", "Hành động", "Đầu vào", "Đầu ra", "Cổng / ghi chú"]
rows3 = []
names = {u[0]: u[1] for u in D.UCS}
for uid, steps in D.FLOWS.items():
    for i, (lane, act, inp, out, note) in enumerate(steps, 1):
        rows3.append([uid, names[uid], i, lane, act, inp, out, note])
ws3 = sheet("3. Luồng chính", H3, rows3, [10, 30, 6, 6, 70, 18, 20, 16], "LuongChinh", freeze="C2")
gate_fill = PatternFill("solid", fgColor=LIGHT)
for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
    lane = row[3].value; note = str(row[7].value or "")
    if lane == "N" or note.startswith("G"):
        for c in row:
            c.fill = gate_fill

# ---------------- 4. Luồng thay thế
H4 = ["Mã UC", "Tên use case", "Mã luồng", "Điều kiện", "Xử lý"]
rows4 = [[uid, names[uid], a, cond, handle] for uid, alts in D.ALTS.items() for (a, cond, handle) in alts]
sheet("4. Luồng thay thế", H4, rows4, [10, 30, 9, 40, 80], "LuongThayThe", freeze="C2")

# ---------------- 5. Kịch bản mẫu
H5 = ["Bước", "Việc xảy ra (kịch bản: board mới, zip hỗn hợp, thiếu thông tin, mô phỏng trước khi có board)", "Use case", "Ai", "Cổng người", "Kết quả / tri thức sinh ra"]
sheet("5. Kịch bản mẫu", H5, [list(s) for s in D.SCENARIO], [6, 70, 30, 22, 18, 44], "KichBanMau")

# ---------------- 6. Truy vết + thống kê bằng công thức
ws6 = wb.create_sheet("6. Truy vết")
ws6.append(["Mã UC", "Tên", "Nhóm", "Ưu tiên", "Mốc", "Cổng người", "FR", "UR", "Màn hình"])
for c in ws6[1]:
    c.font = hdr_font; c.fill = hdr_fill; c.border = border
for u in D.UCS:
    ws6.append([u[0], u[1], u[2], u[11], u[12], u[9], u[13], u[14], u[15]])
for row in ws6.iter_rows(min_row=2, max_row=ws6.max_row):
    for c in row:
        c.font = body_font; c.alignment = wrap; c.border = border
n = ws6.max_row
for i, w in enumerate([10, 34, 22, 8, 7, 16, 24, 22, 22], 1):
    ws6.column_dimensions[get_column_letter(i)].width = w
ws6.freeze_panes = "A2"
# thống kê
r0 = n + 3
ws6.cell(row=r0, column=1, value="THỐNG KÊ (công thức)").font = Font(name=FONT, bold=True, size=12, color=RED)
ws6.cell(row=r0 + 1, column=1, value="Tổng số UC"); ws6.cell(row=r0 + 1, column=2, value=f"=COUNTA(A2:A{n})")
ws6.cell(row=r0 + 2, column=1, value="UC bắt buộc (M)"); ws6.cell(row=r0 + 2, column=2, value=f'=COUNTIF(D2:D{n},"M")')
ws6.cell(row=r0 + 3, column=1, value="UC nên có (S)"); ws6.cell(row=r0 + 3, column=2, value=f'=COUNTIF(D2:D{n},"S")')
ws6.cell(row=r0 + 4, column=1, value="UC có cổng người"); ws6.cell(row=r0 + 4, column=2, value=f'=COUNTIF(F2:F{n},"<>—")')
ws6.cell(row=r0 + 5, column=1, value="Theo nhóm →").font = Font(name=FONT, bold=True, size=10)
groups = sorted({u[2] for u in D.UCS})
for j, g in enumerate(groups):
    ws6.cell(row=r0 + 6 + j, column=1, value=g); ws6.cell(row=r0 + 6 + j, column=2, value=f'=COUNTIF(C2:C{n},"{g}")')
hdr_m = r0 + 7 + len(groups)
ws6.cell(row=hdr_m, column=1, value="Theo mốc →").font = Font(name=FONT, bold=True, size=10)
ws6.cell(row=hdr_m, column=2, value="Số UC").font = Font(name=FONT, bold=True, size=10)
ws6.cell(row=hdr_m, column=3, value="Số bước luồng chính").font = Font(name=FONT, bold=True, size=10)
for j, m in enumerate(["M0", "M1", "M2", "M3", "M4", "M5"]):
    rr = hdr_m + 1 + j
    ws6.cell(row=rr, column=1, value=m); ws6.cell(row=rr, column=2, value=f'=COUNTIF(E2:E{n},"{m}")')
    # số bước luồng chính của mốc: SUMIF trên sheet 2 cột Q (Số bước) theo cột M (Mốc)
    ws6.cell(row=rr, column=3, value=f"=SUMIF('2. Danh mục UC'!M:M,\"{m}\",'2. Danh mục UC'!Q:Q)")
for row in ws6.iter_rows(min_row=r0, max_row=ws6.max_row, max_col=3):
    for c in row:
        if c.font != hdr_font and not c.font.bold:
            c.font = body_font

wb.save("EIDE_Use_Case_Chi_Tiet.xlsx")
print("ok", len(D.UCS), len(rows3), len(rows4))
