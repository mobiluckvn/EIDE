"""`report.export` và `report.human_ai_matrix` — CDS-12.5 REPORT-02, REPORT-03.

`report.export` là chỗ DUY NHẤT dựng docx/pdf (DEV-070), nên test ở đây canh hai thứ mà không
nơi nào khác canh được: thiếu công cụ thì báo TRƯỚC khi tốn một lượt gọi mô hình, và mục Nguồn
được kiểm trên chính tệp sắp giao chứ không tin bước trước.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from eide_core import store
from eide_core.errors import EideError
from eide_core.ledger import Ledger
from eide_core.policy import PolicyGate
from eide_core.router import Context, Router


@pytest.fixture
def du_an(tmp_path, workspace):
    r = Router(gate=PolicyGate(), ledger=Ledger(tmp_path / "l.jsonl"))
    res = r.invoke("project.create", {"text": "báo cáo"}, Context(project_dir=workspace)).result
    root = workspace / res["project_id"]
    store.migrate(store.store_path(root), ledger=r.ledger)
    return r, Context(project_dir=root, extra={"gate": PolicyGate(), "ledger": r.ledger}), root


def _gia_doc(monkeypatch, root: Path, noi_dung: str) -> Path:
    """Thay `doc.generate` bằng một tệp Markdown dựng sẵn — test này không kiểm bộ sinh tài
    liệu (đã có test riêng), nó kiểm bước XUẤT."""
    f = root / ".eide" / "docs" / "bc.md"
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_text(noi_dung, encoding="utf-8")
    import eide.caps.doc as m
    monkeypatch.setattr(m, "generate",
                        lambda p, c: {"doc_id": "doc_x", "path": str(f), "style_issues": 0})
    return f


# ---------------------------------------------------------------- REPORT-02 export


def test_xuat_md_giu_nguyen_va_tra_duong_dan_tuong_doi(du_an, monkeypatch):
    from eide.caps.report import export

    _, ctx, root = du_an
    _gia_doc(monkeypatch, root, "# Báo cáo\n\nNội dung.\n\n## Nguồn\n\n- src_1\n")
    kq = export({"type": "md", "scope": "toàn dự án"}, ctx)
    assert kq["file"].endswith("bc.md") and not Path(kq["file"]).is_absolute()


def test_thieu_cong_cu_thi_E4001_TRUOC_khi_sinh_tai_lieu(du_an, monkeypatch):
    """Để `pandoc` tự báo "command not found" thì người dùng đã trả tiền cho một lượt gọi mô
    hình rồi mới biết là không xuất được."""
    import eide.caps.doc as m
    from eide.caps.report import export

    _, ctx, _ = du_an
    goi = []
    monkeypatch.setattr(m, "generate", lambda p, c: goi.append(1) or {})
    monkeypatch.setattr("eide_core.tools.which", lambda t: None)

    with pytest.raises(EideError) as e:
        export({"type": "docx", "scope": "x"}, ctx)
    assert e.value.code == "E4001" and e.value.data["alternative"] == "type=md"
    assert goi == [], "không được gọi doc.generate khi đã biết là thiếu công cụ"


def test_khong_co_muc_NGUON_thi_khong_xuat(du_an, monkeypatch):
    """Mục Nguồn là thứ phân biệt một báo cáo truy nguyên được với một tờ giấy."""
    from eide.caps.report import export

    _, ctx, root = du_an
    _gia_doc(monkeypatch, root, "# Báo cáo\n\nKhông có mục nguồn nào.\n")
    with pytest.raises(EideError) as e:
        export({"type": "md", "scope": "x"}, ctx)
    assert e.value.code == "E5002" and "Nguồn" in str(e.value)


def test_muc_nguon_nhan_ca_khong_dau(du_an, monkeypatch):
    """Tài liệu bản tiếng Anh (`doc.translate`) ghi "Nguon"/"Source" — phép kiểm không được
    chặt tới mức từ chối chính bản dịch của mình."""
    from eide.caps.report import export

    _, ctx, root = du_an
    _gia_doc(monkeypatch, root, "# BC\n\nx\n\n### Nguon truy nguoc\n\n- src_1\n")
    assert export({"type": "md", "scope": "x"}, ctx)["file"].endswith("bc.md")


# ---------------------------------------------------------------- REPORT-03 human_ai_matrix


def _run(root: Path, cap: str, actor: str, i: int) -> None:
    with store.open_store(store.store_path(root)) as c:
        c.execute("INSERT INTO capability_run (id, cap, args_hash, actor, status, started_at)"
                  " VALUES (?,?,?,?,?,?)",
                  (f"cr_{i:016d}", cap, "h", actor, "done", "2026-09-12T00:00:00Z"))
        c.commit()


def test_tong_cac_o_bang_so_luot_goi(du_an):
    """tc REPORT-03 nguyên văn: "Tổng khớp thống kê". Lệch nghĩa là một pha rơi khỏi bảng."""
    from openpyxl import load_workbook

    from eide.caps.report import human_ai_matrix

    _, ctx, root = du_an
    for i, (cap, actor) in enumerate([
            ("extract.svd", "orchestrator"), ("kg.add_fact", "orchestrator"),
            ("code.build", "orchestrator"), ("code.merge", "human"),
            ("doc.generate", "orchestrator")]):
        _run(root, cap, actor, i)

    f = root / human_ai_matrix({}, ctx)["file"]
    ws = load_workbook(f).active
    hang = [list(r) for r in ws.iter_rows(values_only=True)]
    tong = next(r for r in hang if r and r[0] == "TỔNG")
    assert tong[4] == 5, hang
    assert tong[1] + tong[2] == 5, "mỗi lượt gọi thuộc đúng một trong hai cột AI/người"


def test_namespace_la_roi_vao_Khac_chu_khong_bien_mat(du_an):
    """Một namespace chưa có trong bảng pha phải rơi vào "Khác" — biến mất thì tổng lệch mà
    không ai biết vì sao."""
    from openpyxl import load_workbook

    from eide.caps.report import human_ai_matrix

    _, ctx, root = du_an
    _run(root, "user.tu_che", "orchestrator", 0)
    f = root / human_ai_matrix({}, ctx)["file"]
    pha = [r[0] for r in load_workbook(f).active.iter_rows(values_only=True)]
    assert "Khác" in pha


def test_cot_tri_thuc_dem_dung_nhom_cham_fact(du_an):
    from openpyxl import load_workbook

    from eide.caps.report import human_ai_matrix

    _, ctx, root = du_an
    _run(root, "extract.svd", "orchestrator", 0)
    _run(root, "code.build", "orchestrator", 1)
    f = root / human_ai_matrix({}, ctx)["file"]
    tong = next(r for r in load_workbook(f).active.iter_rows(values_only=True)
                if r and r[0] == "TỔNG")
    assert tong[3] == 1, "chỉ extract.* chạm tri thức, code.build thì không"


def test_ca_bon_nang_luc_report_da_gan_hien_thuc():
    from eide.cli import main  # noqa: F401
    from eide_core.registry import get_registry

    reg = get_registry()
    assert {c.spec.id for c in reg.list(ns="report", implemented=True)} == {
        "report.progress", "report.explain", "report.export", "report.human_ai_matrix"}
