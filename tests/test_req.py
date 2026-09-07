"""Nhóm req.* — CDS-12.1; DDD-14 §2 Requirement.

tc của hợp đồng: REQ-03 TC-67 "ADC 1 MSPS vs fact 2,4 MSPS → ok; yêu cầu 5 MSPS → không ok";
REQ-04 TC-68; REQ-05 "SAFETY luôn M"; REQ-02 "Mỗi yêu cầu có mã duy nhất"; REQ-06 "lỗ hổng liệt
kê"; REQ-07 "Mỗi yêu cầu ≥ 1 AC".

Năm năng lực deterministic kiểm được không cần mô hình. Ba năng lực có sinh (`elicit`,
`classify`, `acceptance`) chỉ kiểm phần mã nguồn làm — gán mã, ghi store, chặn thiếu AC — vì
đó mới là phần có bất biến.
"""
from __future__ import annotations

import json

import pytest

from eide.caps.req import TU_MO_HO, do_duoc
from eide_core import store
from eide_core.ledger import Ledger
from eide_core.policy import PolicyGate
from eide_core.router import Context, Router


@pytest.fixture
def du_an(tmp_path, workspace):
    r = Router(gate=PolicyGate(), ledger=Ledger(tmp_path / "l.jsonl"))
    res = r.invoke("project.create", {"text": "dự án yêu cầu"}, Context(project_dir=workspace)).result
    root = workspace / res["project_id"]
    store.migrate(store.store_path(root), ledger=r.ledger)
    ctx = Context(project_dir=root, extra={"gate": PolicyGate(), "ledger": r.ledger})
    return r, ctx, root


def them_req(root, ds):
    """Ghi thẳng vào store — các test deterministic không nên phụ thuộc `req.classify` (có sinh)."""
    from eide.caps.req import _ghi_requirement
    _ghi_requirement(root, ds)
    return [d["id"] for d in ds]


# ---------- chuẩn hóa đơn vị (nền của REQ-03 và REQ-04)


@pytest.mark.parametrize(("cau", "gia_tri", "dai_luong"), [
    ("bus I2C chạy 400 kHz", 400e3, "tần số"),
    ("bus I2C chạy 0,4 MHz", 400e3, "tần số"),          # dấu phẩy thập phân tiếng Việt
    ("phản hồi ≤ 50 ms", 50e-3, "thời gian"),
    ("ADC 1 MSPS", 1e6, "tốc độ lấy mẫu"),
    ("dòng ngủ 12 uA", 12e-6, "dòng"),
    ("không có số nào ở đây", None, None),
])
def test_do_duoc_chuan_hoa_ve_don_vi_goc(cau, gia_tri, dai_luong):
    """"400 kHz" và "0,4 MHz" phải RA CÙNG MỘT SỐ. Nếu không, REQ-04 báo chúng mâu thuẫn — một
    cảnh báo giả, và cảnh báo giả lặp lại làm người ta thôi đọc cảnh báo thật."""
    kq = do_duoc(cau)
    if gia_tri is None:
        assert kq is None
    else:
        assert kq[1] == dai_luong
        assert kq[0] == pytest.approx(gia_tri)


def test_hai_cach_viet_cung_nguong_khong_bi_bao_mau_thuan(du_an):
    """Hệ quả trực tiếp của test trên, kiểm ở mức năng lực."""
    r, ctx, root = du_an
    ids = them_req(root, [
        {"id": "FR-COM-01", "kind": "FR", "text": "Bus I2C của cảm biến chạy ở 400 kHz"},
        {"id": "FR-COM-02", "kind": "FR", "text": "Bus I2C của cảm biến chạy ở 0,4 MHz"},
    ])
    issues = r.invoke("req.detect_conflict", {"reqset_ids": ids}, ctx).result["issues"]
    assert [i for i in issues if i["kind"] == "conflict"] == []


# ---------- REQ-04 req.detect_conflict (deterministic)


def test_cung_dai_luong_nguong_trai_nhau_la_conflict(du_an):
    """tc TC-68. Quy tắc 1 của REQ-04."""
    r, ctx, root = du_an
    ids = them_req(root, [
        {"id": "FR-COM-01", "kind": "FR", "text": "Bus I2C cảm biến chạy ở 100 kHz"},
        {"id": "FR-COM-02", "kind": "FR", "text": "Bus I2C cảm biến chạy ở 400 kHz"},
    ])
    xd = [i for i in r.invoke("req.detect_conflict", {"reqset_ids": ids}, ctx).result["issues"]
          if i["kind"] == "conflict"]
    assert len(xd) == 1
    assert set(xd[0]["req_ids"]) == {"FR-COM-01", "FR-COM-02"}
    assert xd[0]["suggestion"]


def test_hai_bus_khac_nhau_khong_phai_mau_thuan(du_an):
    """"I2C 400 kHz" và "SPI 8 MHz" đều là tần số nhưng nói về hai thứ khác nhau. Không có
    `_cung_chu_de`, mọi cặp yêu cầu cùng đơn vị đều thành mâu thuẫn và bảng issues thành rác."""
    r, ctx, root = du_an
    ids = them_req(root, [
        {"id": "FR-COM-01", "kind": "FR", "text": "Bus I2C cảm biến chạy 400 kHz"},
        {"id": "FR-COM-02", "kind": "FR", "text": "Bus SPI thẻ nhớ chạy 8 MHz"},
    ])
    issues = r.invoke("req.detect_conflict", {"reqset_ids": ids}, ctx).result["issues"]
    assert [i for i in issues if i["kind"] == "conflict"] == []


def test_tu_mo_ho_bi_bat_va_neu_ten(du_an):
    """Quy tắc 3. Nêu ĐÍCH DANH từ mơ hồ, vì "câu này mơ hồ" thì người sửa vẫn phải đoán chỗ nào."""
    r, ctx, root = du_an
    ids = them_req(root, [{"id": "NFR-01", "kind": "NFR",
                           "text": "Hệ thống phải phản hồi nhanh và ổn định"}])
    mh = [i for i in r.invoke("req.detect_conflict", {"reqset_ids": ids}, ctx).result["issues"]
          if i["kind"] == "ambiguous"]
    assert len(mh) == 1
    assert set(mh[0]["tu_mo_ho"]) == {"nhanh", "ổn định"}


def test_cau_do_duoc_khong_bi_bao_mo_ho(du_an):
    """Ngược lại của test trên — nếu quy tắc bắt cả câu đã có ngưỡng thì nó vô dụng."""
    r, ctx, root = du_an
    ids = them_req(root, [{"id": "NFR-01", "kind": "NFR",
                           "text": "Hệ thống phản hồi trong 50 ms kể từ khi nhấn nút"}])
    assert r.invoke("req.detect_conflict", {"reqset_ids": ids}, ctx).result["issues"] == []


def test_yeu_cau_NFR_khong_co_nguong_la_unmeasurable(du_an):
    """Quy tắc 2. Chỉ áp cho NFR/RT/HW: một FR như "hiển thị nhiệt độ lên LCD" không cần con số."""
    r, ctx, root = du_an
    ids = them_req(root, [
        {"id": "NFR-01", "kind": "NFR", "text": "Bộ nhớ chương trình phải vừa vi điều khiển"},
        {"id": "FR-UI-01", "kind": "FR", "text": "Hiển thị nhiệt độ lên màn hình"},
    ])
    kq = r.invoke("req.detect_conflict", {"reqset_ids": ids}, ctx).result["issues"]
    assert [i["req_ids"] for i in kq if i["kind"] == "unmeasurable"] == [["NFR-01"]]


def test_moi_tu_mo_ho_trong_bang_deu_bat_duoc():
    """Bảng TU_MO_HO là QUY TẮC, không phải gợi ý cho mô hình — nên nó phải thực sự có tác dụng.
    Một từ nằm trong bảng mà biểu thức chính quy không khớp là quy tắc chết."""
    import re
    for t in TU_MO_HO:
        assert re.search(rf"\b{t}\b", f"hệ thống phải {t} hơn".lower()), t


# ---------- REQ-03 req.ground_hw (TC-67)


def them_fact(root, subject, value, unit, fid="f_0000000000000001"):
    with store.open_store(store.store_path(root)) as c:
        c.execute("INSERT OR REPLACE INTO source (id, uri, sha256, kind, tier, fetched_at) "
                  "VALUES ('s_ds','ds.pdf','x','datasheet','gold','2026-01-01T00:00:00Z')")
        c.execute("INSERT OR REPLACE INTO fact (id, subject, predicate, value, unit, source_id,"
                  " method, tier, confidence, status) VALUES (?,?,?,?,?,?,?,?,?,?)",
                  (fid, subject, "timing", json.dumps(value), unit, "s_ds", "parser",
                   "gold", 1.0, "verified"))
        c.commit()


def test_TC67_yeu_cau_duoi_kha_nang_thi_ok(du_an):
    """TC-67 vế 1: "ADC 1 MSPS vs fact 2,4 MSPS → ok"."""
    r, ctx, root = du_an
    them_fact(root, "chip:st.stm32f411ce/periph:ADC1", 2.4, "MSPS")
    ids = them_req(root, [{"id": "UR-SNS-01", "kind": "HW",
                           "text": "ADC lấy mẫu cảm biến ở 1 MSPS"}])
    rep = r.invoke("req.ground_hw",
                   {"reqset_ids": ids, "passport": "st.stm32f411ce@1.2.0"}, ctx).result["report"]
    assert rep[0]["ok"] is True
    assert rep[0]["facts"] == ["f_0000000000000001"]


def test_TC67_yeu_cau_vuot_kha_nang_thi_khong_ok_va_co_de_xuat(du_an):
    """TC-67 vế 2: "yêu cầu 5 MSPS → không ok". Bước 2 của hợp đồng đòi đề xuất thay thế — báo
    "không khả thi" mà không nói làm gì tiếp thì người dùng vẫn kẹt."""
    r, ctx, root = du_an
    them_fact(root, "chip:st.stm32f411ce/periph:ADC1", 2.4, "MSPS")
    ids = them_req(root, [{"id": "UR-SNS-01", "kind": "HW",
                           "text": "ADC lấy mẫu cảm biến ở 5 MSPS"}])
    rep = r.invoke("req.ground_hw",
                   {"reqset_ids": ids, "passport": "st.stm32f411ce@1.2.0"}, ctx).result["report"]
    assert rep[0]["ok"] is False
    assert rep[0]["alternative"]


def test_khong_co_fact_thi_chua_xac_dinh_chu_khong_doan(du_an):
    """Bước 1: "không có fact → 'chưa xác định' + kg.request". `ok=None` chứ KHÔNG phải False,
    và tuyệt đối không phải True: một yêu cầu được đánh "khả thi" bằng phỏng đoán sẽ đi tiếp vào
    kiến trúc, vào mã, và chỉ lộ ra khi board đã đặt về."""
    r, ctx, root = du_an
    ids = them_req(root, [{"id": "UR-SNS-01", "kind": "HW", "text": "ADC lấy mẫu ở 1 MSPS"}])
    rep = r.invoke("req.ground_hw",
                   {"reqset_ids": ids, "passport": "st.stm32f411ce@1.2.0"}, ctx).result["report"]
    assert rep[0]["ok"] is None
    assert "kg.request" in rep[0]["alternative"]


def test_fact_cua_chip_khac_khong_duoc_dung_de_ket_luan(du_an):
    """Lọc theo hộ chiếu. Fact của chip khác dùng để kết luận "khả thi" cho board này còn tệ hơn
    không có fact nào — nó sai một cách tự tin."""
    r, ctx, root = du_an
    them_fact(root, "chip:nxp.imxrt1062/periph:ADC1", 2.4, "MSPS")
    ids = them_req(root, [{"id": "UR-SNS-01", "kind": "HW", "text": "ADC lấy mẫu ở 1 MSPS"}])
    rep = r.invoke("req.ground_hw",
                   {"reqset_ids": ids, "passport": "st.stm32f411ce@1.2.0"}, ctx).result["report"]
    assert rep[0]["ok"] is None


def test_ground_hw_ghi_feasibility_vao_store(du_an):
    """DDD-14: cột `feasibility` là JSON {ok, facts[], note}. Kết luận phải ở lại store, không
    chỉ ở giá trị trả về — `req.trace_matrix` và `arch.*` sau này đọc từ đó."""
    r, ctx, root = du_an
    them_fact(root, "chip:st.stm32f411ce/periph:ADC1", 2.4, "MSPS")
    ids = them_req(root, [{"id": "UR-SNS-01", "kind": "HW", "text": "ADC lấy mẫu ở 1 MSPS"}])
    r.invoke("req.ground_hw", {"reqset_ids": ids, "passport": "st.stm32f411ce@1.2.0"}, ctx)
    with store.open_store(store.store_path(root)) as c:
        (fe,) = c.execute("SELECT feasibility FROM requirement WHERE id='UR-SNS-01'").fetchone()
    assert json.loads(fe)["ok"] is True


# ---------- REQ-05 req.prioritize


def test_SAFETY_luon_M(du_an):
    """tc: "SAFETY luôn M". Quy tắc CỨNG ở mã: một yêu cầu an toàn bị hạ xuống "Should" là yêu
    cầu sẽ bị cắt khi hết thời gian — đúng lúc người ta cần nó nhất."""
    r, ctx, root = du_an
    ids = them_req(root, [
        {"id": "UR-SAF-01", "kind": "SAFETY", "text": "Dừng động cơ trong 100 ms khi mất tín hiệu",
         "priority": "W"},
        {"id": "NFR-01", "kind": "NFR", "text": "Khởi động trong 2 s"},
    ])
    out = r.invoke("req.prioritize", {"reqset_ids": ids}, ctx).result["reqset"]
    uu = {x["id"]: x["priority"] for x in out}
    assert uu["UR-SAF-01"] == "M"          # ghi đè cả "W" đã có sẵn
    assert uu["NFR-01"] == "C"


def test_SAFETY_giu_M_ca_khi_khong_kha_thi_tren_board(du_an):
    """Đây mới là bất biến thật của "SAFETY luôn M" — test trên nó chỉ trùng với bảng tra.

    `_uu_tien_theo` hạ một bậc khi `feasibility.ok is False` (chưa làm được thì chưa thể là
    Must). Áp quy tắc ấy cho SAFETY thì một yêu cầu an toàn khó làm trên board này sẽ tự tụt
    xuống S — đúng cái mà quy tắc cứng phải chặn, và là lý do nó đứng TRƯỚC mọi điều chỉnh khác.
    """
    r, ctx, root = du_an
    them_req(root, [{"id": "UR-SAF-01", "kind": "SAFETY",
                     "text": "Dừng động cơ trong 100 ms khi mất tín hiệu",
                     "feasibility": {"ok": False, "facts": [], "note": "vượt khả năng"}}])
    out = r.invoke("req.prioritize", {"reqset_ids": ["UR-SAF-01"]}, ctx).result["reqset"]
    assert out[0]["priority"] == "M"


def test_khong_kha_thi_thi_ha_mot_bac_voi_yeu_cau_thuong(du_an):
    """Vế đối chứng: nếu không có đường nào hạ bậc thì test trên xanh vì lý do sai."""
    r, ctx, root = du_an
    them_req(root, [{"id": "FR-CTL-01", "kind": "FR", "text": "Điều khiển PID 100 Hz",
                     "feasibility": {"ok": False, "facts": [], "note": "vượt khả năng"}}])
    out = r.invoke("req.prioritize", {"reqset_ids": ["FR-CTL-01"]}, ctx).result["reqset"]
    assert out[0]["priority"] == "C"          # S hạ một bậc


def test_tac_tu_khong_duoc_ha_uu_tien_M(du_an):
    """ask "Đổi ưu tiên M↔S" ở mức T1*. Gán lần đầu thì tự làm được; hạ một yêu cầu ĐÃ là M
    xuống thấp hơn là quyết định phạm vi — việc của người."""
    r, ctx, root = du_an
    ids = them_req(root, [{"id": "FR-CTL-01", "kind": "FR", "text": "Điều khiển PID 100 Hz",
                           "priority": "M"}])
    res = r.invoke("req.prioritize", {"reqset_ids": ids}, ctx)
    assert res.status == "failed"
    assert res.error["eide_code"] == "E3000"
    with store.open_store(store.store_path(root)) as c:
        (p,) = c.execute("SELECT priority FROM requirement WHERE id='FR-CTL-01'").fetchone()
    assert p == "M", "bị chặn thì KHÔNG được ghi giá trị mới xuống store"


def test_nguoi_thi_ha_duoc_uu_tien_M(du_an):
    """Vế còn lại: chặn tác tử mà cũng chặn luôn người thì năng lực này vô dụng."""
    r, _, root = du_an
    ctx = Context(project_dir=root, actor="human",
                  extra={"gate": PolicyGate(), "ledger": r.ledger})
    ids = them_req(root, [{"id": "FR-CTL-01", "kind": "FR", "text": "Điều khiển PID 100 Hz",
                           "priority": "M"}])
    out = r.invoke("req.prioritize", {"reqset_ids": ids}, ctx).result["reqset"]
    assert out[0]["priority"] == "S"


# ---------- REQ-06 req.trace_matrix


def test_liet_ke_lo_hong_thieu_module_va_test(du_an):
    """tc: "lỗ hổng liệt kê". Giá trị nằm ở cột gaps, không ở ma trận."""
    r, ctx, root = du_an
    them_req(root, [
        {"id": "FR-SNS-01", "kind": "FR", "text": "Đọc nhiệt độ mỗi 1 s",
         "trace": ["mod_sensor", "TC-10"]},
        {"id": "FR-CTL-01", "kind": "FR", "text": "Điều khiển PID 100 Hz", "trace": []},
    ])
    out = r.invoke("req.trace_matrix", {}, ctx).result
    thieu = {(g["req_id"], g["thieu"]) for g in out["gaps"]}
    assert thieu == {("FR-CTL-01", "module"), ("FR-CTL-01", "test")}
    assert out["file"].endswith(".md")


def test_ma_tran_khong_cho_loc_tap_con(du_an):
    """Hợp đồng KHÔNG khai `reqset_ids`. Cho lọc thì sẽ có người xuất ma trận "không lỗ hổng"
    bằng cách chọn đúng những yêu cầu đã xong."""
    from pathlib import Path
    ct = json.loads(Path("docs/spec/cds.json").read_text(encoding="utf-8"))
    sc = next(c for c in ct if c["id"] == "req.trace_matrix")["input_schema"]
    assert "reqset_ids" not in sc["properties"]


def test_xuat_xlsx_khi_duoc_yeu_cau(du_an):
    """`format` có hai giá trị. Trả về .md khi người dùng xin .xlsx là nói dối về đầu ra."""
    r, ctx, root = du_an
    them_req(root, [{"id": "FR-SNS-01", "kind": "FR", "text": "Đọc nhiệt độ mỗi 1 s"}])
    f = r.invoke("req.trace_matrix", {"format": "xlsx"}, ctx).result["file"]
    assert f.endswith(".xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(f)
    assert wb["Truy vết"].cell(row=2, column=1).value == "FR-SNS-01"
    assert wb["Lỗ hổng"].max_row == 3          # tiêu đề + thiếu module + thiếu test


# ---------- REQ-08 req.change_impact


def test_tac_dong_tu_3_module_thi_hoi_nguoi(du_an):
    """ask "Tác động > ngưỡng (≥ 3 module)". Ba module trở lên không còn là sửa cục bộ mà là đổi
    thiết kế."""
    r, ctx, root = du_an
    them_req(root, [{"id": "FR-CTL-01", "kind": "FR", "text": "PID 100 Hz",
                     "trace": ["mod_a", "mod_b", "mod_c", "TC-1"]}])
    res = r.invoke("req.change_impact", {"delta": {"req_ids": ["FR-CTL-01"]}}, ctx)
    assert res.status == "failed"
    assert res.error["eide_code"] == "E3000"


def test_tac_dong_hai_module_thi_lam_thang(du_an):
    """Ngưỡng phải có VẾ CHO QUA, nếu không thì mọi thay đổi đều hỏi và người sẽ bấm bừa."""
    r, ctx, root = du_an
    them_req(root, [{"id": "FR-CTL-01", "kind": "FR", "text": "PID 100 Hz",
                     "trace": ["mod_a", "mod_b", "TC-1"]}])
    im = r.invoke("req.change_impact", {"delta": {"req_ids": ["FR-CTL-01"]}}, ctx).result["impact"]
    assert im["modules"] == ["mod_a", "mod_b"]
    assert im["tests"] == ["TC-1"]


# ---------- REQ-02 req.classify: phần mã nguồn làm (không phải mô hình)


def test_ma_yeu_cau_khong_bao_gio_lap(du_an, monkeypatch):
    """tc: "Mỗi yêu cầu có mã duy nhất". Mã do MÃ NGUỒN gán: gọi hai lần phải ra hai dải mã
    khác nhau, kể cả khi mô hình trả về y hệt."""
    r, ctx, root = du_an
    _gia_lap(monkeypatch, {"reqset": [
        {"kind": "FR", "text": "Đọc cảm biến nhiệt độ mỗi 1 s"},
        {"kind": "SAFETY", "text": "Watchdog dừng an toàn trong 100 ms"},
        {"kind": "NFR", "text": "Khởi động trong 2 s"},
    ]})
    raw = [{"text": x} for x in ("a", "b", "c")]
    m1 = [x["id"] for x in r.invoke("req.classify", {"raw": raw}, ctx).result["reqset"]]
    m2 = [x["id"] for x in r.invoke("req.classify", {"raw": raw}, ctx).result["reqset"]]
    assert m1 == ["FR-SNS-01", "UR-SAF-01", "NFR-01"]
    assert m2 == ["FR-SNS-02", "UR-SAF-02", "NFR-02"]
    assert not set(m1) & set(m2)


def test_NFR_khong_mang_ma_nhom(du_an, monkeypatch):
    """DDD-14 §2 cho đúng ba dạng: `UR-xx-nn | FR-xx-nn | NFR-nn`. Một yêu cầu phi chức năng
    ("khởi động ≤ 2 s") cắt ngang nhiều nhóm — gán nó vào SNS hay CTL đều sai."""
    r, ctx, _ = du_an
    _gia_lap(monkeypatch, {"reqset": [
        {"kind": "NFR", "text": "Cảm biến khởi động trong 2 s"}]})   # có từ khóa SNS
    assert r.invoke("req.classify", {"raw": [{"text": "x"}]},
                    ctx).result["reqset"][0]["id"] == "NFR-01"


def test_classify_ghi_store_voi_status_generated(du_an, monkeypatch):
    """Bước 2: "Ghi bảng requirement status generated". `generated` chứ không `accepted` —
    yêu cầu do mô hình viết ra chưa ai duyệt."""
    r, ctx, root = du_an
    _gia_lap(monkeypatch, {"reqset": [{"kind": "FR", "text": "Đọc cảm biến mỗi 1 s"}]})
    r.invoke("req.classify", {"raw": [{"text": "x"}]}, ctx)
    with store.open_store(store.store_path(root)) as c:
        assert c.execute("SELECT status FROM requirement").fetchone()[0] == "generated"


# ---------- REQ-01 / REQ-07: bất biến do mã nguồn giữ


def test_elicit_khong_co_dau_vao_thi_bao_E5002(du_an):
    """ask "Ô trống không có mặc định". Ba ô đều tùy chọn, nhưng trống hết thì không có gì để
    tách — trả về `raw: []` im lặng sẽ khiến chuỗi sau tưởng dự án không có yêu cầu nào."""
    r, ctx, _ = du_an
    res = r.invoke("req.elicit", {}, ctx)
    assert res.status == "failed"
    assert res.error["eide_code"] == "E5002"


def test_elicit_giu_locator_cho_moi_yeu_cau(du_an, monkeypatch):
    """tc: "≥ 15 yêu cầu thô CÓ LOCATOR". Không có nó thì không ai trả lời được "yêu cầu này từ
    đâu ra", và một yêu cầu không truy nguyên được thì không tranh luận được."""
    r, ctx, _ = du_an
    _gia_lap(monkeypatch, {"raw": [{"text": f"yêu cầu {i}"} for i in range(15)]})
    raw = r.invoke("req.elicit", {"text": "mô tả robot", "sources": ["README.md"]},
                   ctx).result["raw"]
    assert len(raw) == 15
    assert all(x["locator"] is not None and x["source"] == "README.md" for x in raw)


def test_acceptance_thieu_mot_yeu_cau_thi_bao_loi(du_an, monkeypatch):
    """tc: "Mỗi yêu cầu ≥ 1 AC". Mô hình bỏ sót một yêu cầu là chuyện thường; nhận im lặng thì
    yêu cầu ấy sẽ không có cách nào kiểm, và `trace_matrix` báo lỗ hổng ở tận cuối chuỗi."""
    r, ctx, root = du_an
    ids = them_req(root, [
        {"id": "FR-SNS-01", "kind": "FR", "text": "Đọc nhiệt độ mỗi 1 s"},
        {"id": "FR-CTL-01", "kind": "FR", "text": "PID 100 Hz"},
    ])
    _gia_lap(monkeypatch, {"acceptance": [
        {"req_id": "FR-SNS-01", "given": "g", "when": "w", "then": "t",
         "observable": "serial_pattern"}]})
    res = r.invoke("req.acceptance", {"reqset_ids": ids}, ctx)
    assert res.status == "failed"
    assert res.error["eide_code"] == "E5002"
    assert "FR-CTL-01" in str(res.error)


def _gia_lap(monkeypatch, data):
    """Thay Gateway bằng câu trả lời cố định — các test này kiểm phần MÃ NGUỒN làm quanh mô hình."""
    class _R:
        def __init__(self, d): self.data = d
    class _G:
        def run(self, *a, **k): return _R(data)
    import eide.caps.req as m
    monkeypatch.setattr(m, "_gateway", lambda ctx: _G())
    monkeypatch.setattr(m, "_ngu_canh", lambda ctx, t: "")
